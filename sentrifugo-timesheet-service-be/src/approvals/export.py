from __future__ import annotations

import io
from typing import Any

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _build_rows(detail: dict[str, Any]) -> tuple[list[str], list[list[Any]]]:
    headers = ["Project", "Task", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun", "Total"]
    rows: list[list[Any]] = []
    for row in detail.get("weekly_timeline") or []:
        rows.append([
            row.get("project_name") or row.get("project_id", ""),
            row.get("task_name") or row.get("task_id", ""),
            row.get("mon", 0),
            row.get("tue", 0),
            row.get("wed", 0),
            row.get("thu", 0),
            row.get("fri", 0),
            row.get("sat", 0),
            row.get("sun", 0),
            row.get("total", 0),
        ])
    return headers, rows


def _write_week_sheet(ws, detail: dict[str, Any]) -> None:
    """Render one week's timesheet (the standard per-week layout) into ``ws``."""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    user_name = detail.get("user_name") or detail.get("user_id", "")
    ws.merge_cells("A1:J1")
    ws["A1"] = f"Timesheet - {user_name}"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "Employee ID:"
    ws["B2"] = detail.get("emp_code") or detail.get("user_id", "")
    ws["A3"] = "Week:"
    week_start = str(detail.get("week_start_date", ""))[:10]
    week_end = str(detail.get("week_end_date", ""))[:10]
    ws["B3"] = f"{week_start} to {week_end}"
    ws["A4"] = "Status:"
    ws["B4"] = str(detail.get("timesheet_status", "")).replace("_", " ").title()
    ws["A5"] = "Total Hours:"
    ws["B5"] = detail.get("total_hours", 0)

    for cell in ["A2", "A3", "A4", "A5"]:
        ws[cell].font = Font(bold=True)
    ws["B5"].font = Font(bold=True)

    headers, rows = _build_rows(detail)
    start_row = 7

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row_data in enumerate(rows, start_row + 1):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx >= 3:
                cell.alignment = Alignment(horizontal="center")

    grand_total_row = start_row + len(rows) + 1
    total_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    for col_idx in range(1, 11):
        cell = ws.cell(row=grand_total_row, column=col_idx)
        cell.fill = total_fill
        cell.border = thin_border
    gt_label = ws.cell(row=grand_total_row, column=1, value="Grand Total")
    gt_label.font = Font(bold=True)
    gt_hours = ws.cell(row=grand_total_row, column=10, value=detail.get("total_hours", 0))
    gt_hours.font = Font(bold=True)
    gt_hours.alignment = Alignment(horizontal="center")

    if detail.get("approval_history"):
        ah_start = grand_total_row + 2
        ws.cell(row=ah_start, column=1, value="History").font = Font(bold=True, size=11)

        ah_headers = ["Date & Time", "Role", "Action", "Comments"]
        ah_header_row = ah_start + 1
        for col_idx, h in enumerate(ah_headers, 1):
            cell = ws.cell(row=ah_header_row, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            cell.alignment = Alignment(horizontal="center")

        for row_idx, ar in enumerate(detail["approval_history"], ah_header_row + 1):
            acted_at = str(ar.get("acted_at", ""))[:19]
            role = str(ar.get("approver_role", "")).replace("_", " ").title()
            action = str(ar.get("action", "")).replace("_", " ").title()
            comments = ar.get("comments") or ""
            ws.cell(row=row_idx, column=1, value=acted_at).border = thin_border
            ws.cell(row=row_idx, column=2, value=role).border = thin_border
            ws.cell(row=row_idx, column=3, value=action).border = thin_border
            ws.cell(row=row_idx, column=4, value=comments).border = thin_border

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 20
    for col_letter in "CDEFGHIJ":
        ws.column_dimensions[col_letter].width = 8


def generate_excel(detail: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Timesheet"
    _write_week_sheet(ws, detail)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _safe_sheet_title(base: str, used: set[str]) -> str:
    """Sanitise/uniquify an Excel sheet title (<=31 chars, no : \\ / ? * [ ])."""
    invalid = set(r':\/?*[]')
    title = "".join(c for c in base if c not in invalid).strip()[:31] or "Sheet"
    candidate = title
    i = 2
    while candidate in used:
        suffix = f" ({i})"
        candidate = f"{title[:31 - len(suffix)]}{suffix}"
        i += 1
    used.add(candidate)
    return candidate


def _write_summary_sheet(ws, month: dict[str, Any]) -> None:
    """Render the month overview (totals + one row per week) into ``ws``."""
    header_font = Font(bold=True, size=11)
    header_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    user_name = month.get("user_name") or month.get("user_id", "")
    ws.merge_cells("A1:D1")
    ws["A1"] = f"Monthly Timesheet - {user_name}"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "Employee ID:"
    ws["B2"] = month.get("emp_code") or month.get("user_id", "")
    ws["A3"] = "Period:"
    ws["B3"] = f"{str(month.get('period_start', ''))[:10]} to {str(month.get('period_end', ''))[:10]}"
    ws["A4"] = "Status:"
    ws["B4"] = str(month.get("timesheet_status", "")).replace("_", " ").title()
    ws["A5"] = "Total Hours:"
    ws["B5"] = month.get("total_hours", 0)
    ws["A6"] = "Weeks:"
    ws["B6"] = month.get("week_count", len(month.get("weeks", [])))
    for cell in ["A2", "A3", "A4", "A5", "A6"]:
        ws[cell].font = Font(bold=True)

    weeks = month.get("weeks", [])
    start_row = 8
    headers = ["Week Start", "Week End", "Total Hours", "Status"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row_idx, wk in enumerate(weeks, start_row + 1):
        ws.cell(row=row_idx, column=1, value=str(wk.get("week_start_date", ""))[:10]).border = thin_border
        ws.cell(row=row_idx, column=2, value=str(wk.get("week_end_date", ""))[:10]).border = thin_border
        hrs = ws.cell(row=row_idx, column=3, value=wk.get("total_hours", 0))
        hrs.border = thin_border
        hrs.alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=str(wk.get("timesheet_status", "")).replace("_", " ").title()).border = thin_border

    gt_row = start_row + 1 + len(weeks)
    ws.cell(row=gt_row, column=2, value="Grand Total").font = Font(bold=True)
    gt = ws.cell(row=gt_row, column=3, value=month.get("total_hours", 0))
    gt.font = Font(bold=True)
    gt.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 18


def generate_monthly_excel(month: dict[str, Any]) -> bytes:
    """Workbook with a Summary sheet plus one sheet per week."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    _write_summary_sheet(summary, month)

    used = {"Summary"}
    for idx, wk in enumerate(month.get("weeks", []), 1):
        base = str(wk.get("week_start_date", ""))[:10]
        title = _safe_sheet_title(f"Wk {base}" if base else f"Week {idx}", used)
        ws = wb.create_sheet(title=title)
        _write_week_sheet(ws, wk)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


STATUS_LABELS = {
    "l1_approved": "Pending Review",
    "client_approved": "Approved",
    "client_rejected": "Rejected",
}


def generate_review_excel(items: list[dict[str, Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Client Review"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="6B21A8", end_color="6B21A8", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.merge_cells("A1:F1")
    ws["A1"] = "Timesheets - Client Review Export"
    ws["A1"].font = Font(bold=True, size=14)

    headers = ["Employee", "Project", "Week Start", "Week End", "Total Hours", "Status"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row_idx, item in enumerate(items, 4):
        ws.cell(row=row_idx, column=1, value=item.get("user_name") or item.get("user_id", "")).border = thin_border
        ws.cell(row=row_idx, column=2, value=item.get("project_name", "")).border = thin_border
        ws_date = str(item.get("week_start_date", ""))[:10]
        we_date = str(item.get("week_end_date", ""))[:10]
        ws.cell(row=row_idx, column=3, value=ws_date).border = thin_border
        ws.cell(row=row_idx, column=4, value=we_date).border = thin_border
        hrs_cell = ws.cell(row=row_idx, column=5, value=item.get("total_hours", 0))
        hrs_cell.border = thin_border
        hrs_cell.alignment = Alignment(horizontal="center")
        status = item.get("timesheet_status", "")
        ws.cell(row=row_idx, column=6, value=STATUS_LABELS.get(status, status)).border = thin_border

    total_row = len(items) + 4
    ws.cell(row=total_row, column=4, value="Grand Total").font = Font(bold=True)
    total_hrs = sum(item.get("total_hours", 0) for item in items)
    total_cell = ws.cell(row=total_row, column=5, value=total_hrs)
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 25
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_activity_excel(
    items: list[dict[str, Any]],
    *,
    title: str = "Activity History Export",
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Activity History"

    header_font = Font(bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="6B21A8", end_color="6B21A8", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    ws.merge_cells("A1:H1")
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)

    headers = ["Date & Time", "Decision Maker", "Role", "Employee", "Projects", "Hours", "Action", "Comments"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for row_idx, item in enumerate(items, 4):
        acted = item.get("acted_at")
        acted_str = str(acted)[:19] if acted else ""
        ws.cell(row=row_idx, column=1, value=acted_str).border = thin_border
        ws.cell(row=row_idx, column=2, value=item.get("approver_name", "")).border = thin_border
        ws.cell(row=row_idx, column=3, value=str(item.get("approver_role", "")).replace("_", " ").title()).border = thin_border
        ws.cell(row=row_idx, column=4, value=item.get("employee_name", "")).border = thin_border
        ws.cell(row=row_idx, column=5, value=item.get("project_names", "")).border = thin_border
        hrs_cell = ws.cell(row=row_idx, column=6, value=item.get("hours", 0))
        hrs_cell.border = thin_border
        hrs_cell.alignment = Alignment(horizontal="center")
        action = str(item.get("action", "")).replace("_", " ").title()
        ws.cell(row=row_idx, column=7, value=action).border = thin_border
        ws.cell(row=row_idx, column=8, value=item.get("comments", "")).border = thin_border

    total_row = len(items) + 4
    ws.cell(row=total_row, column=5, value="Total").font = Font(bold=True)
    total_hrs = sum(item.get("hours", 0) for item in items)
    total_cell = ws.cell(row=total_row, column=6, value=total_hrs)
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["F"].width = 10
    ws.column_dimensions["G"].width = 12
    ws.column_dimensions["H"].width = 35

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _write_week_page(pdf: FPDF, detail: dict[str, Any]) -> None:
    """Render one week's timesheet onto the current PDF page."""
    user_name = detail.get("user_name") or detail.get("user_id", "")
    week_start = str(detail.get("week_start_date", ""))[:10]
    week_end = str(detail.get("week_end_date", ""))[:10]

    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, f"Timesheet - {user_name}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10)
    pdf.cell(35, 6, "Employee ID:", new_x="END")
    pdf.cell(0, 6, detail.get("emp_code") or detail.get("user_id", ""), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Week:", new_x="END")
    pdf.cell(0, 6, f"{week_start} to {week_end}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Status:", new_x="END")
    pdf.cell(0, 6, str(detail.get("timesheet_status", "")).replace("_", " ").title(), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Total Hours:", new_x="END")
    pdf.cell(0, 6, str(detail.get("total_hours", 0)), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    headers, rows = _build_rows(detail)
    col_widths = [40, 30, 14, 14, 14, 14, 14, 14, 14, 16]

    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(226, 239, 218)
    for i, h in enumerate(headers):
        pdf.cell(col_widths[i], 8, h, border=1, fill=True, align="C")
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    line_h = 7
    for row_data in rows:
        x0 = pdf.get_x()
        y0 = pdf.get_y()

        # Determine row height from wrapping text columns (Project & Task)
        proj_lines = pdf.multi_cell(col_widths[0], line_h, str(row_data[0]), split_only=True)
        task_lines = pdf.multi_cell(col_widths[1], line_h, str(row_data[1]), split_only=True)
        row_height = max(len(proj_lines), len(task_lines), 1) * line_h

        # Project column – render text then draw full-height border rect
        pdf.set_xy(x0, y0)
        pdf.multi_cell(col_widths[0], line_h, str(row_data[0]), border=0, align="L")
        pdf.rect(x0, y0, col_widths[0], row_height)

        # Task column
        task_x = x0 + col_widths[0]
        pdf.set_xy(task_x, y0)
        pdf.multi_cell(col_widths[1], line_h, str(row_data[1]), border=0, align="L")
        pdf.rect(task_x, y0, col_widths[1], row_height)

        # Numeric columns – single cell spanning full row height
        cur_x = task_x + col_widths[1]
        for i, val in enumerate(row_data[2:], 2):
            pdf.set_xy(cur_x, y0)
            pdf.cell(col_widths[i], row_height, str(val), border=1, align="C")
            cur_x += col_widths[i]

        # Advance past this row
        pdf.set_xy(x0, y0 + row_height)

    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(col_widths[0], 8, "Grand Total", border=1)
    for i in range(1, 9):
        pdf.cell(col_widths[i], 8, "", border=1)
    pdf.cell(col_widths[9], 8, str(detail.get("total_hours", 0)), border=1, align="C")
    pdf.ln()

    if detail.get("approval_history"):
        pdf.ln(8)
        pdf.set_font("Helvetica", "B", 11)
        pdf.cell(0, 8, "History", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 9)
        for ar in detail["approval_history"]:
            action = ar.get("action", "")
            role = ar.get("approver_role", "")
            acted_at = str(ar.get("acted_at", ""))[:19]
            comments = ar.get("comments") or ""
            pdf.cell(0, 6, f"{acted_at} - {role} - {action} {f'({comments})' if comments else ''}", new_x="LMARGIN", new_y="NEXT")


def generate_pdf(detail: dict[str, Any]) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    _write_week_page(pdf, detail)
    return bytes(pdf.output())


def _write_summary_page(pdf: FPDF, month: dict[str, Any]) -> None:
    """Render the month overview (totals + one row per week) onto the current page."""
    user_name = month.get("user_name") or month.get("user_id", "")
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, f"Monthly Timesheet - {user_name}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10)
    pdf.cell(35, 6, "Employee ID:", new_x="END")
    pdf.cell(0, 6, str(month.get("emp_code") or month.get("user_id", "")), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Period:", new_x="END")
    pdf.cell(0, 6, f"{str(month.get('period_start', ''))[:10]} to {str(month.get('period_end', ''))[:10]}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Status:", new_x="END")
    pdf.cell(0, 6, str(month.get("timesheet_status", "")).replace("_", " ").title(), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(35, 6, "Total Hours:", new_x="END")
    pdf.cell(0, 6, str(month.get("total_hours", 0)), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    headers = ["Week Start", "Week End", "Hours", "Status"]
    widths = [40, 40, 25, 50]
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(226, 239, 218)
    for i, h in enumerate(headers):
        pdf.cell(widths[i], 8, h, border=1, fill=True, align="C")
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for wk in month.get("weeks", []):
        pdf.cell(widths[0], 7, str(wk.get("week_start_date", ""))[:10], border=1)
        pdf.cell(widths[1], 7, str(wk.get("week_end_date", ""))[:10], border=1)
        pdf.cell(widths[2], 7, str(wk.get("total_hours", 0)), border=1, align="C")
        pdf.cell(widths[3], 7, str(wk.get("timesheet_status", "")).replace("_", " ").title(), border=1)
        pdf.ln()

    pdf.set_font("Helvetica", "B", 9)
    pdf.cell(widths[0] + widths[1], 8, "Grand Total", border=1)
    pdf.cell(widths[2], 8, str(month.get("total_hours", 0)), border=1, align="C")
    pdf.cell(widths[3], 8, "", border=1)
    pdf.ln()


def generate_monthly_pdf(month: dict[str, Any]) -> bytes:
    """PDF with a summary page plus one page per week."""
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    _write_summary_page(pdf, month)
    for wk in month.get("weeks", []):
        pdf.add_page()
        _write_week_page(pdf, wk)
    return bytes(pdf.output())
