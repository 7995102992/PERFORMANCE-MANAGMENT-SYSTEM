"""Shared helpers for building import/export spreadsheets.

All import templates (clients, projects, tasks) use the same visual format:
a styled header row, " *" markers + a darker fill for required columns,
auto-sized widths, a frozen header, optional per-column format hints, and
in-cell dropdowns for enumerated fields. Keeping the styling here guarantees
every template looks and parses identically.
"""

from __future__ import annotations

from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="366092")
REQUIRED_FILL = PatternFill("solid", fgColor="244062")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_header_row(
    ws: Worksheet,
    columns: Sequence[str],
    *,
    required_columns: Iterable[str] = (),
    hints: dict[str, str] | None = None,
) -> None:
    """Write the standard styled header row into row 1 of ``ws``.

    Required columns get a " *" suffix and a darker fill; all columns are
    auto-sized, the header row is frozen, and any provided ``hints`` are
    attached as cell comments (format tooltips).
    """
    required = set(required_columns)
    hints = hints or {}
    for col_idx, label in enumerate(columns, start=1):
        is_required = label in required
        display = f"{label} *" if is_required else label
        cell = ws.cell(row=1, column=col_idx, value=display)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if is_required else HEADER_FILL
        cell.alignment = HEADER_ALIGN
        ws.column_dimensions[get_column_letter(col_idx)].width = max(len(display) + 4, 18)
        hint = hints.get(label)
        if hint:
            cell.comment = Comment(f"Format: {hint}", "System")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 28


def parse_header(cells: Iterable) -> list[str]:
    """Normalise a header row to lowercase keys, stripping the " *" required marker."""
    return [str(c).replace("*", "").strip().lower() if c else "" for c in cells]


def col_letter_for(columns: Sequence[str], label: str) -> str:
    """Excel column letter (1-based) of ``label`` within ``columns``."""
    return get_column_letter(list(columns).index(label) + 1)


def add_list_dropdown(
    ws: Worksheet,
    column_letter: str,
    options: Sequence[str],
    *,
    last_row: int = 1000,
    allow_blank: bool = True,
    error: str = "Pick a value from the dropdown.",
    prompt: str | None = None,
    prompt_title: str | None = None,
) -> None:
    """Attach an in-cell list dropdown to ``column_letter`` for data rows 2..last_row.

    ``options`` are embedded inline, so each value must be short and free of
    commas (Excel uses the comma as the item separator for inline lists).
    """
    formula = '"' + ",".join(options) + '"'
    dv = DataValidation(
        type="list",
        formula1=formula,
        allow_blank=allow_blank,
        showErrorMessage=True,
    )
    dv.error = error
    dv.errorTitle = "Invalid value"
    if prompt:
        dv.prompt = prompt
        dv.promptTitle = prompt_title or "Allowed values"
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{last_row}")


def format_date_column(
    ws: Worksheet,
    column_letter: str,
    *,
    last_row: int = 1000,
    number_format: str = "yyyy-mm-dd",
    prompt: str = "Enter a date — it is stored as YYYY-MM-DD.",
) -> None:
    """Format a column as ISO dates and validate that entries are real dates.

    Pre-formatting the column makes Excel store typed input as a true date value
    (which the importer reads back as an ISO date) instead of locale-specific
    text like ``01-06-2026`` that the strict ``YYYY-MM-DD`` parser would reject.
    """
    ws.column_dimensions[column_letter].number_format = number_format
    dv = DataValidation(
        type="date",
        operator="greaterThanOrEqual",
        formula1="DATE(1900,1,1)",
        allow_blank=True,
        showErrorMessage=True,
    )
    dv.error = "Enter a valid date (it will be stored as YYYY-MM-DD)."
    dv.errorTitle = "Invalid date"
    dv.prompt = prompt
    dv.promptTitle = "Date"
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{last_row}")


def add_sheet_dropdown(
    wb: Workbook,
    ws: Worksheet,
    column_letter: str,
    options: Sequence[str],
    *,
    ref_sheet: str = "_REF",
    ref_col: int = 1,
    last_row: int = 1000,
    allow_blank: bool = True,
    error: str = "Pick a value from the dropdown.",
) -> None:
    """Attach a dropdown whose options live in a hidden helper sheet.

    Use this instead of :func:`add_list_dropdown` when the option list is long
    or user-supplied (e.g. client names) — Excel caps inline list formulas at
    ~255 characters, but a range reference has no such limit.
    """
    if not options:
        return
    ref = wb[ref_sheet] if ref_sheet in wb.sheetnames else wb.create_sheet(ref_sheet)
    ref.sheet_state = "hidden"
    src_col = get_column_letter(ref_col)
    for i, val in enumerate(options, start=1):
        ref.cell(row=i, column=ref_col, value=val)
    dv = DataValidation(
        type="list",
        formula1=f"='{ref_sheet}'!${src_col}$1:${src_col}${len(options)}",
        allow_blank=allow_blank,
        showErrorMessage=True,
    )
    dv.error = error
    dv.errorTitle = "Invalid value"
    ws.add_data_validation(dv)
    dv.add(f"{column_letter}2:{column_letter}{last_row}")
