"""Import ONLY the intern / contractual (consultant) employees from the
`Employees_data_import-template` sheet into an existing org.

The template sheet holds the whole workforce (Permanent + Consultant + Intern).
Full-time/permanent people are already in the DB, so this script imports just
the two contingent populations and leaves everything else untouched:

    EmploymentType 'Intern' / 'Apprentice'              -> employment_type internship, emp_code SIL-I-<num>
    EmploymentType 'Consultant' / '* Contract'          -> employment_type contract,   emp_code SIL-C-<num>
    anything else (Permanent, Probation, blank, ...)    -> SKIPPED (reported, never written)

INCREMENTAL and non-destructive, like an "add missing people" pass:
  * a row whose emp_code OR email is already in the DB -> reported as EXISTS, not
    inserted, nothing updated;
  * only genuinely new rows are inserted (user + employee doc);
  * the same person twice (same email, two codes — the sheet has several consultants
    with an earlier and a later contract) keeps BOTH employee stints but only ONE
    login: the latest stint owns the user, earlier ones are flagged `is_rehire`.
    Same when the email already exists in the DB (`--skip-existing-email` to skip
    those rows instead);
  * the BU is REUSED (never created) — the org must already exist and have one.
    Its emp-code counters are advanced past the imported C/I numbers so the app's
    generator resumes after them;
  * departments / designations are reused by name and created only when missing.

Active rule: the sheet is COLOUR CODED — a row written in red is a person who has
exited, any other row is a current employee. That colour is authoritative:

    red row      -> EXIT  (employment_status=exit, date_of_exit = DateOfLeaving if
                    the sheet has one, user INACTIVE, never a password and never
                    an activation email)
    normal row   -> ACTIVE (DateOfLeaving, when present, is just the CONTRACT END
                    date — it is reported as `contract_ends`, never as date_of_exit)

`DateOfLeaving` alone is not a reliable signal here (several current contractors
carry a future end date, and a few red rows carry none), so it is only used as the
fallback when the source has no colours at all — a .csv export, say — where a row
counts as active when its leaving date is missing or still ahead of the as-of date
(default today, override --as-of). Rows where the colour and the date disagree are
listed in the report.

employment_status for active rows: contract -> direct-contract, intern -> internship
(the master-data entry is created if the deployment does not have it yet).

Account state for ACTIVE rows:
    default          -> user ACTIVE with password test@123 (already-activated import)
    --pending        -> user INACTIVE, no password (activate later)
    --activation     -> user INACTIVE + activation email enqueued
                        (outbox -> RabbitMQ -> schedule service)

Report (always written, dry run AND commit):
    scripts/intern_contract_import_report.md
    scripts/intern_contract_import_report.csv

Usage:
    cd Sentrifugo-IAM-Admin-BE
    python -m scripts.import_intern_contract --org-id <oid>                     # DRY RUN
    python -m scripts.import_intern_contract --org-id <oid> --commit            # import
    python -m scripts.import_intern_contract --org-id <oid> --commit --activation
    python -m scripts.import_intern_contract --org-id <oid> --commit --emit-events
    python -m scripts.import_intern_contract --org-id <oid> --revert            # undo this import
    python -m scripts.import_intern_contract --file "<path .xlsx/.csv>"         # other source file

On Windows run with PYTHONIOENCODING=utf-8 to avoid cp1252 console errors.
"""

import argparse
import asyncio
import csv
import re
import sys
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from beanie import PydanticObjectId, init_beanie
from bson import ObjectId
from motor.motor_asyncio import AsyncIOMotorClient

from src.auth.models import UserDocument
from src.auth.utils.tools import get_password_hash
from src.config import settings
from src.correlation import get_correlation_id
from src.master_data.models import MasterDataDocument
from src.models import OutboxEventDocument, StatusEnum
from src.rabbitmq import outbox
from src.modules.organisation.models import (
    AddressDocument,
    BusinessUnitDocument,
    DepartmentDocument,
    DesignationDocument,
    EmployeeDocument,
    OrganisationDocument,
)
from scripts.datafile import _cell, find_file, read_rows

# ── Config ────────────────────────────────────────────────────────────────
REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_FILE = REPO_ROOT / "data-dir" / "Employees_data_import-template (2).xlsx"
ORG_ADMIN_EMAIL = "admin+sil@sagarsoft.in"
ORG_ID_FALLBACK = "6a4c8cae5457ae1842d39985"
MARKER = "import_intern_contract"   # every doc we create carries this in created_by
DEFAULT_PASSWORD = "test@123"
DEFAULT_PREFIX = "SIL"
PAD = 4

NOW = datetime.now(timezone.utc)
REPORT_DIR = Path(__file__).resolve().parent
REPORT_MD = REPORT_DIR / "intern_contract_import_report.md"
REPORT_CSV = REPORT_DIR / "intern_contract_import_report.csv"

# EmploymentType (or EmployeeStatus, as a fallback) -> which population a row is in.
INTERN_TYPES = {"intern", "interns", "internship", "apprentice", "trainee"}
CONTRACT_TYPES = {"consultant", "contract", "contractual", "contractor",
                  "direct contract", "third party contract"}

# employment_status key used for ACTIVE rows, per population.
ACTIVE_STATUS_KEY = {"contract": "direct-contract", "intern": "internship"}
STATUS_VALUE = {"direct-contract": "Direct Contract", "internship": "Internship"}

# Junk / non-employee rows dropped outright.
SKIP_IDS = {"1234", "SIL-099999", "BGCK273", "admin", "SIL-01"}


# ── Helpers ───────────────────────────────────────────────────────────────
def norm_name(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").strip()).lower()


def sorted_key(s: str) -> str:
    """Token-order-insensitive key so 'Arsid Srinivas' matches 'Srinivas Arsid'."""
    return " ".join(sorted(norm_name(s).split()))


def norm_dept(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").replace("&", "and").strip())


def code_number(emp_id: str) -> str:
    """Exact numeric suffix as written (leading zeros kept). '' if none."""
    m = re.match(r"^[A-Za-z]*[-]?(\d+)$", (emp_id or "").strip())
    return m.group(1) if m else ""


def parse_date(value):
    """Any date representation -> a `date` (or None). Handles real datetime cells,
    strings in any layout, and raw Excel serial numbers (this sheet mixes them)."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return _excel_serial(value)
    s = str(value).strip()
    if not s or s.lower() in ("none", "null", "-"):
        return None
    if re.fullmatch(r"\d{5}", s):           # e.g. '46387' -> 2026-12-31
        return _excel_serial(int(s))
    # ISO first (the sheet stores '2023-07-04 00:00:00' as text) — dayfirst would
    # flip those into 4 July -> 7 April.
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    try:
        from dateutil import parser as _du
        return _du.parse(s, dayfirst=True).date()
    except Exception:
        return None


def _excel_serial(n):
    try:
        n = float(n)
    except (TypeError, ValueError):
        return None
    if not (20000 <= n <= 80000):           # outside ~1954..2119: not a date serial
        return None
    return (datetime(1899, 12, 30) + timedelta(days=n)).date()


def parse_dt(value):
    d = parse_date(value)
    return datetime(d.year, d.month, d.day, tzinfo=timezone.utc) if d else None


def cell(row: dict, *names) -> str:
    for n in names:
        v = row.get(n)
        if v is not None and str(v).strip() and str(v).strip().lower() != "none":
            return str(v).strip()
    return ""


# ── Colour reading (the sheet marks exited people in red) ─────────────────
RED_KEY = "__red"


def _rgb_is_red(rgb) -> bool:
    """True for any strong red (FFFF0000, FFC00000, ...)."""
    if not isinstance(rgb, str) or len(rgb) < 6:
        return False
    try:
        r, g, b = (int(rgb[-6:][i:i + 2], 16) for i in (0, 2, 4))
    except ValueError:
        return False
    return r >= 0x90 and g <= 0x66 and b <= 0x66


def _cell_is_red(c) -> bool:
    col = getattr(c.font, "color", None)
    if col is not None and getattr(col, "type", None) == "rgb" and _rgb_is_red(col.rgb):
        return True
    fill = c.fill
    if fill is not None and fill.patternType == "solid":
        fg = fill.fgColor
        if getattr(fg, "type", None) == "rgb" and _rgb_is_red(fg.rgb):
            return True
    return False


def read_rows_colored(path: Path, sheet: str | None = None) -> tuple[list[dict], bool]:
    """Rows exactly as `scripts.datafile.read_rows` gives them, each carrying a
    `__red` flag: True when the row is colour coded red (= the person exited).

    Returns (rows, has_colours). A .csv carries no formatting, so has_colours is
    False there and the caller falls back to the date rule. A row counts as red
    only when at least half of its filled cells are red — a single highlighted
    cell (the sheet flags odd MaritalStatus values that way) is not an exit mark."""
    if path.suffix.lower() not in (".xlsx", ".xlsm"):
        return read_rows(path), False

    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)          # styles need a non-read-only load
    if sheet and sheet in wb.sheetnames:
        ws = wb[sheet]
    elif "Data" in wb.sheetnames:
        ws = wb["Data"]
    else:
        ws = wb[wb.sheetnames[0]]

    it = ws.iter_rows()
    header = [str(c.value).strip() if c.value is not None else "" for c in next(it)]
    rows, any_red = [], False
    for cells in it:
        if all(c.value is None for c in cells):
            continue
        row = {header[i]: _cell(cells[i].value if i < len(cells) else None)
               for i in range(len(header))}
        filled = [c for c in cells[:len(header)] if c.value not in (None, "")]
        reds = [c for c in filled if _cell_is_red(c)]
        row[RED_KEY] = bool(filled) and len(reds) * 2 >= len(filled)
        any_red = any_red or row[RED_KEY]
        rows.append(row)
    wb.close()
    return rows, any_red


def audit(created_on=None, modified_on=None):
    return {"created_by": MARKER, "created_on": created_on or NOW,
            "modified_by": MARKER, "modified_on": modified_on or NOW}


# ── Parse phase (pure; no DB) ─────────────────────────────────────────────
def classify(etype: str, estat: str) -> str | None:
    """'intern' | 'contract' | None (= not part of this import)."""
    for value in (etype, estat):
        v = norm_name(value)
        if not v:
            continue
        if v in INTERN_TYPES or "intern" in v or "apprentice" in v:
            return "intern"
        if v in CONTRACT_TYPES or "consultant" in v or "contract" in v:
            return "contract"
    return None


def parse_rows(rows: list[dict], prefix: str, as_of: date | None = None,
               use_colour: bool = True) -> tuple[list[dict], dict]:
    """Classify every sheet row. Returns (records, max numbers per code letter).

    Red row = exited. `as_of` (default today) is only the fallback for sources
    without colours: DateOfLeaving is the contract END date, so a date still ahead
    of as_of means the person is on contract, not gone."""
    as_of = as_of or date.today()
    records: list[dict] = []
    maxes = {"C": 0, "I": 0}

    for r in rows:
        raw_id = cell(r, "employeeId", "EmployeeId", "employee_id")
        etype = cell(r, "EmploymentType")
        estat = cell(r, "EmployeeStatus")
        dept = norm_dept(cell(r, "Department"))
        rec = {
            "raw_id": raw_id,
            "email": cell(r, "EmailAddress", "Email", "WorkEmail").lower(),
            "first": cell(r, "FirstName"),
            "last": cell(r, "LastName"),
            "etype": etype, "estat": estat,
            "gender": cell(r, "Gender"),
            "marital": cell(r, "MaritalStatus"),
            "designation": cell(r, "Designation"),
            "department": dept,
            "doj": parse_date(r.get("DateOfJoining")),
            "dol": parse_date(r.get("DateOfLeaving")),
            "mgr1": cell(r, "ReportingManager"),
            "mgr2": cell(r, "L2ReportingManager", "l2ReportingManager"),
            "created_on": parse_dt(r.get("createddate")) or NOW,
            "modified_on": parse_dt(r.get("modifieddate")) or NOW,
            "red": r.get(RED_KEY),          # True = colour coded as exited
            "kind": None, "emp_code": "", "note": "",
        }

        if not raw_id and not rec["email"]:
            continue                                    # blank spacer row
        kind = classify(etype, estat)
        if kind is None:
            rec.update(action="skip",
                       reason=f"not intern/contract (EmploymentType='{etype or '-'}')")
            records.append(rec); continue
        rec["kind"] = kind

        if raw_id in SKIP_IDS:
            rec.update(action="skip", reason="junk / admin row")
            records.append(rec); continue
        if not rec["email"]:
            rec.update(action="skip", reason="no email")
            records.append(rec); continue

        num = code_number(raw_id)
        if not num:
            rec.update(action="skip", reason=f"employee id '{raw_id}' has no number")
            records.append(rec); continue

        letter = "I" if kind == "intern" else "C"
        rec["emp_code"] = f"{prefix}-{letter}-{num}"
        rec["num"] = num
        # Red = exited. Without colours fall back to the contract end date: still
        # ahead (or absent) => the engagement is running.
        by_date = rec["dol"] is None or rec["dol"] > as_of
        rec["active"] = (not rec["red"]) if use_colour else by_date
        rec["rule"] = "colour" if use_colour else "date"
        rec["disagrees"] = use_colour and rec["red"] is not None and by_date != rec["active"]
        # An active person's leaving date is just when the contract runs out.
        rec["ends_on"] = rec["dol"] if (rec["active"] and rec["dol"]) else None
        rec.update(action="insert", reason="",
                   user_status="active" if rec["active"] else "exit")
        try:
            maxes[letter] = max(maxes[letter], int(num))
        except ValueError:
            pass
        records.append(rec)

    # Same email twice inside the sheet = a rehired person with two stints (the
    # sheet has several: an earlier consultant contract and a later one). We keep
    # BOTH employee records but only ONE login: the most recent stint (latest DoJ)
    # is the primary and owns the user; earlier stints are flagged is_rehire and
    # point at it. Email is a unique login id, so no second account is made.
    by_email: dict[str, dict] = {}
    for r in sorted([x for x in records if x["action"] == "insert"],
                    key=lambda x: (x["doj"] or date.min), reverse=True):
        primary = by_email.get(r["email"])
        if primary:
            r["rehire_of"] = primary["emp_code"]
            r["note"] = f"earlier stint of {primary['emp_code']} — same login, flagged is_rehire"
        else:
            by_email[r["email"]] = r
            r["rehire_of"] = None
    return records, maxes


# ── DB helpers ────────────────────────────────────────────────────────────
async def _md_id(category: str, key: str):
    doc = await MasterDataDocument.find_one(
        MasterDataDocument.category == category,
        MasterDataDocument.key == key,
        MasterDataDocument.organisation_id == None,  # noqa: E711
    )
    if not doc:
        doc = await MasterDataDocument.find_one(
            MasterDataDocument.category == category, MasterDataDocument.key == key)
    return doc.id if doc else None


async def _ensure_status(key: str, value: str):
    existing = await MasterDataDocument.find_one(
        MasterDataDocument.category == "EMPLOYMENT_STATUSES",
        MasterDataDocument.key == key)
    if existing:
        return existing.id
    doc = MasterDataDocument(category="EMPLOYMENT_STATUSES", key=key, value=value,
                             organisation_id=None, is_active=True, is_custom=False)
    await doc.insert()
    print(f"  + master data EMPLOYMENT_STATUSES/{key} created")
    return doc.id


async def _md_map(category: str) -> dict:
    out: dict[str, PydanticObjectId] = {}
    for d in await MasterDataDocument.find(MasterDataDocument.category == category).to_list():
        out[(d.value or "").strip().lower()] = d.id
        out[(d.key or "").strip().lower()] = d.id
    return out


async def pick_bu(org_id: PydanticObjectId, prefix: str):
    """REUSE the org's business unit — never create one here."""
    bus = await BusinessUnitDocument.find(
        BusinessUnitDocument.organisation_id == org_id,
        BusinessUnitDocument.deleted_on == None,  # noqa: E711
    ).to_list()
    if not bus:
        raise RuntimeError("org has no business unit — run the main employee import first")
    bu = next((b for b in bus if (b.emp_code_prefix or "").upper() == prefix.upper()), None)
    return bu or bus[0]


async def mark_existing(records: list[dict], org_id: PydanticObjectId, prefix: str,
                        skip_existing_email: bool = False):
    """Flag rows already present in the DB.

    emp_code already in the DB            -> EXISTS (never re-inserted, never updated).
    email already belongs to a user       -> that person is already in the system with a
                                             different code, i.e. this sheet row is another
                                             stint: the employee record is inserted and
                                             LINKED to the existing login (flagged is_rehire),
                                             or reported as EXISTS when --skip-existing-email.
    """
    emps = await EmployeeDocument.find(
        EmployeeDocument.organisation_id == org_id,
        EmployeeDocument.deleted_on == None,  # noqa: E711
    ).to_list()
    codes = {(e.emp_code or "").upper() for e in emps}
    numbers = {}
    for e in emps:
        n = code_number((e.emp_code or "").split("-")[-1])
        if n:
            numbers.setdefault(n.lstrip("0") or "0", e.emp_code)
    # Raw motor read: legacy user docs can hold values the current model rejects
    # (e.g. status='exit' from the first migration) — we only need email + org.
    emails = {}
    async for u in UserDocument.get_motor_collection().find(
            {}, {"email": 1, "organisation_id": 1}):
        if u.get("email"):
            emails[u["email"].lower()] = u

    code_by_uid: dict[str, str] = {}
    for e in emps:
        if e.user_id:
            code_by_uid.setdefault(str(e.user_id), e.emp_code or "")

    for r in [x for x in records if x["action"] == "insert"]:
        if r["emp_code"].upper() in codes:
            r.update(action="exists", reason=f"emp_code {r['emp_code']} already in the DB")
            continue
        if r["email"] in emails and not r.get("rehire_of"):
            u = emails[r["email"]]
            uid = u["_id"]
            if str(u.get("organisation_id", "")) != str(org_id):
                r.update(action="exists",
                         reason=f"email belongs to a user in ANOTHER org (user {uid})")
                continue
            if skip_existing_email:
                r.update(action="exists",
                         reason=f"email already belongs to user {uid} (--skip-existing-email)")
                continue
            r["_link_uid"] = PydanticObjectId(uid)
            r["_link_code"] = code_by_uid.get(str(uid), "")
            r["note"] = (f"person already in the DB as {r['_link_code'] or 'a user'} — "
                         f"new stint linked to that login, flagged is_rehire")
            continue
        other = numbers.get((r["num"] or "").lstrip("0") or "0")
        if other:
            r["note"] = f"number also used by {other} (different code — still imported)"


# ── Write phase ───────────────────────────────────────────────────────────
async def cleanup(org_id: PydanticObjectId):
    print(f"Removing anything created by this importer (marker={MARKER})...")
    coll = EmployeeDocument.get_motor_collection()
    ours = [e.emp_code for e in await EmployeeDocument.find(
        {"organisation_id": org_id, "created_by": MARKER}).to_list() if e.emp_code]
    if ours:  # undo the has_prior_stint flags we stamped on pre-existing stints
        await coll.update_many({"organisation_id": org_id, "prior_emp_code": {"$in": ours}},
                               {"$unset": {"has_prior_stint": "", "prior_emp_code": ""}})
    emp = await EmployeeDocument.find({"organisation_id": org_id, "created_by": MARKER}).delete()
    usr = await UserDocument.find({"organisation_id": org_id, "created_by": MARKER}).delete()
    des = await DesignationDocument.find({"organisation_id": org_id, "created_by": MARKER}).delete()
    dep = await DepartmentDocument.find({"organisation_id": org_id, "created_by": MARKER}).delete()
    print(f"  removed: {getattr(emp,'deleted_count',0)} employees, {getattr(usr,'deleted_count',0)} users, "
          f"{getattr(des,'deleted_count',0)} designations, {getattr(dep,'deleted_count',0)} departments")


async def run(records: list[dict], maxes: dict, org_id: PydanticObjectId, bu,
              mode: str = "active", emit_events: bool = False):
    inserts = [r for r in records if r["action"] == "insert"]
    if not inserts:
        print("Nothing new to insert.")
        return set()

    # 1. master data
    et = {"contract": await _md_id("EMPLOYMENT_TYPES", "contract"),
          "intern": await _md_id("EMPLOYMENT_TYPES", "internship")}
    missing = [k for k, v in et.items() if not v]
    if missing:
        raise RuntimeError(f"Missing EMPLOYMENT_TYPES master data: {missing}. Run seed_master_data first.")
    es = {"exit": await _md_id("EMPLOYMENT_STATUSES", "exit")}
    for kind in {r["kind"] for r in inserts if r["active"]}:
        key = ACTIVE_STATUS_KEY[kind]
        es[key] = await _ensure_status(key, STATUS_VALUE[key])
    gmap = await _md_map("GENDERS")
    mmap = await _md_map("MARITAL_STATUSES")

    # 2. designations — reuse by name, create only what is missing
    existing_desg = {(d.designation_name or "").strip().lower(): d.id
                     for d in await DesignationDocument.find(
                         DesignationDocument.organisation_id == org_id,
                         DesignationDocument.deleted_on == None).to_list()}  # noqa: E711
    desg_ids, desg_created = {}, 0
    for name in sorted({r["designation"] for r in inserts if r["designation"]}):
        hit = existing_desg.get(name.strip().lower())
        if hit:
            desg_ids[name] = hit
            continue
        d = DesignationDocument(organisation_id=org_id, designation_name=name,
                                description="", is_active=True, **audit())
        await d.insert()
        desg_ids[name] = d.id
        desg_created += 1
    print(f"  designations: {desg_created} created, {len(desg_ids) - desg_created} reused")

    # 3. departments — reuse by name, create only what is missing (mapped to the BU)
    existing_dept = {(d.department_name or "").strip().lower(): d.id
                     for d in await DepartmentDocument.find(
                         DepartmentDocument.organisation_id == org_id,
                         DepartmentDocument.deleted_on == None).to_list()}  # noqa: E711
    dept_ids, dept_created = {}, 0
    for name in sorted({r["department"] for r in inserts if r["department"]}):
        hit = existing_dept.get(name.strip().lower())
        if hit:
            dept_ids[name] = hit
            continue
        dep = DepartmentDocument(organisation_id=org_id, department_name=name,
                                 business_units=[bu.id], primary_business_unit=bu.id,
                                 is_active=True, **audit())
        await dep.insert()
        dept_ids[name] = dep.id
        dept_created += 1
        if emit_events:
            await outbox.publish("department.created", {
                "correlation_id": get_correlation_id(),
                "department_id": str(dep.id), "organisation_id": str(org_id),
                "name": name, "business_unit_ids": [str(bu.id)],
                "business_unit_id": str(bu.id), "is_active": True,
                "department_head": None, "department_code": None,
            }, idempotency_key=f"department.created:{dep.id}")
    print(f"  departments: {dept_created} created, {len(dept_ids) - dept_created} reused")

    # 4. users + employee docs. Primary stints first, so an earlier stint of the
    # same person finds the login it must link to (no second account is created).
    pending_activation: list[tuple] = []
    primary_of: dict[str, tuple] = {}          # emp_code -> (user_id, emp_code)
    new_users = 0
    for r in sorted(inserts, key=lambda x: bool(x.get("rehire_of"))):
        active = r["active"]
        link = None
        if r.get("_link_uid"):                 # person already in the DB
            link = (r["_link_uid"], r.get("_link_code") or "")
        elif r.get("rehire_of") and r["rehire_of"] in primary_of:
            link = primary_of[r["rehire_of"]]

        if link:
            uid, keeper_code = link
        else:
            # Account status is only ACTIVE/INACTIVE — the exit lifecycle lives on
            # employee.employment_status, so a left employee is simply INACTIVE.
            if not active:
                u_status, u_pw, u_pwchanged = StatusEnum.INACTIVE, None, None
            elif mode == "active":
                u_status, u_pw, u_pwchanged = StatusEnum.ACTIVE, get_password_hash(DEFAULT_PASSWORD), NOW
            else:                              # 'pending' / 'activation'
                u_status, u_pw, u_pwchanged = StatusEnum.INACTIVE, None, None
            u = UserDocument(
                email=r["email"], password_hash=u_pw, auth_method="local",
                first_name=r["first"] or "-", last_name=r["last"] or "",
                gender=gmap.get(r["gender"].lower()), marital_status=mmap.get(r["marital"].lower()),
                status=u_status, organisation_id=org_id, policy_ids=[],
                password_changed_at=u_pwchanged,
                **audit(r["created_on"], r["modified_on"]),
            )
            await u.insert()
            new_users += 1
            uid, keeper_code = u.id, r["emp_code"]
            primary_of[r["emp_code"]] = (u.id, r["emp_code"])
            if active and mode == "activation":
                pending_activation.append((u.id, r["email"], f"{r['first']} {r['last']}".strip()))
        r["_uid"] = uid

        status_id = es[ACTIVE_STATUS_KEY[r["kind"]]] if active else es["exit"]
        emp = EmployeeDocument(
            organisation_id=org_id, user_id=uid, emp_code=r["emp_code"],
            business_unit_id=bu.id, department_id=dept_ids.get(r["department"]),
            designation_id=desg_ids.get(r["designation"]),
            employment_type=et[r["kind"]], employment_status=status_id,
            date_of_joining=r["doj"], date_of_exit=None if active else r["dol"],
            **audit(r["created_on"], r["modified_on"]),
        )
        await emp.insert()
        r["_eid"] = emp.id
        r["_status_id"] = status_id
        if link:
            # EmployeeDocument has no rehire fields — set them raw, exactly like the
            # main importer: flag this stint and mark the kept one as having a prior.
            r["_rehire"] = True
            coll = EmployeeDocument.get_motor_collection()
            await coll.update_one({"_id": emp.id},
                                  {"$set": {"is_rehire": True, "rehire_primary_user_id": uid,
                                            "rehire_primary_emp_code": keeper_code}})
            if keeper_code:
                await coll.update_one(
                    {"organisation_id": org_id, "user_id": uid, "emp_code": keeper_code},
                    {"$set": {"has_prior_stint": True, "prior_emp_code": r["emp_code"]}})
    print(f"  inserted {len(inserts)} employees ({sum(1 for r in inserts if r['active'])} active, "
          f"{sum(1 for r in inserts if not r['active'])} exit); "
          f"{new_users} new logins, {len(inserts) - new_users} linked to an existing account")

    # 5. managers — resolve by name against everyone in the org (old + new)
    name_to_uid: dict[str, list[tuple]] = {}
    skey_to_uid: dict[str, list[tuple]] = {}
    async for u in UserDocument.get_motor_collection().find(
            {"organisation_id": ObjectId(org_id)},
            {"first_name": 1, "last_name": 1, "status": 1}):
        full = f"{u.get('first_name') or ''} {u.get('last_name') or ''}"
        is_active = u.get("status") == StatusEnum.ACTIVE
        name_to_uid.setdefault(norm_name(full), []).append((u["_id"], is_active))
        skey_to_uid.setdefault(sorted_key(full), []).append((u["_id"], is_active))

    def resolve(name: str):
        cands = name_to_uid.get(norm_name(name)) or skey_to_uid.get(sorted_key(name)) or []
        if not cands:
            return None
        actives = [uid for uid, a in cands if a]
        return actives[0] if actives else cands[0][0]

    unresolved = set()
    for r in inserts:
        l1 = resolve(r["mgr1"]) if r["mgr1"] else None
        l2 = resolve(r["mgr2"]) if r["mgr2"] else None
        r["_l1"], r["_l2"] = l1, l2
        if r["mgr1"] and not l1:
            unresolved.add(r["mgr1"])
        if r["mgr2"] and not l2:
            unresolved.add(r["mgr2"])
        if l1 or l2:
            emp = await EmployeeDocument.get(r["_eid"])
            emp.l1_manager_id, emp.l2_manager_id = l1, l2
            await emp.save()
    print(f"  managers resolved; {len(unresolved)} distinct names unresolved")

    # 6. advance the BU emp-code counters past the numbers we just used
    start = dict(bu.emp_code_start_from or {})
    last = dict(bu.emp_code_last_numbers or {})
    pad = dict(bu.emp_code_padding or {})
    for letter in ("C", "I"):
        if not maxes.get(letter):
            continue
        app_next = start.get(letter, 0) + last.get(letter, 0)
        start[letter] = max(maxes[letter] + 1, app_next)
        last[letter] = 0
        pad.setdefault(letter, PAD)
    bu.emp_code_start_from, bu.emp_code_last_numbers, bu.emp_code_padding = start, last, pad
    if not bu.emp_code_prefix:
        bu.emp_code_prefix = DEFAULT_PREFIX
    await bu.save()
    print(f"  BU '{bu.business_unit_name}' counters resume at {start}")

    # 7. domain events for downstream services
    if emit_events:
        for r in inserts:
            await outbox.publish("employee.created", {
                "correlation_id": get_correlation_id(),
                "user_id": str(r["_uid"]), "employee_id": str(r["_eid"]),
                "organisation_id": str(org_id), "emp_code": r["emp_code"],
                "work_email": r["email"], "name": f"{r['first']} {r['last']}".strip(),
                "first_name": r["first"], "last_name": r["last"],
                "l1_manager_id": str(r["_l1"]) if r.get("_l1") else None,
                "l2_manager_id": str(r["_l2"]) if r.get("_l2") else None,
                "designation_id": str(desg_ids[r["designation"]]) if r["designation"] in desg_ids else None,
                "department_id": str(dept_ids[r["department"]]) if r["department"] in dept_ids else None,
                "business_unit_id": str(bu.id),
                "employment_status": str(r["_status_id"]) if r.get("_status_id") else None,
                "project_status": None,
                "date_of_joining": str(r["doj"]) if r["doj"] else None,
            }, idempotency_key=f"employee.created:{r['_uid']}")
        print(f"  emitted {len(inserts)} employee.created event(s) -> outbox")

    # 8. activation emails (activation mode only) — never for exit rows
    if mode == "activation" and pending_activation:
        from src.valkey import init_valkey
        from src.auth.service import send_activation_email
        await init_valkey()
        sent = 0
        for uid, email, name in pending_activation:
            try:
                await send_activation_email(str(uid), email, name, tenant_id=str(org_id))
                sent += 1
            except Exception as exc:
                print(f"  activation email FAILED for {email}: {exc}")
        print(f"  activation enqueued for {sent}/{len(pending_activation)} active employees")

    return unresolved


def dryrun_managers(records, known_names: set, known_skeys: set) -> set:
    unresolved = set()
    for r in records:
        if r["action"] != "insert":
            continue
        for m in (r["mgr1"], r["mgr2"]):
            if m and norm_name(m) not in known_names and sorted_key(m) not in known_skeys:
                unresolved.add(m)
    return unresolved


# ── Report ────────────────────────────────────────────────────────────────
def write_report(records, unresolved, committed: bool, source: Path, mode: str, as_of: date,
                 colour_rule: bool):
    inserts = [r for r in records if r["action"] == "insert"]
    exists = [r for r in records if r["action"] == "exists"]
    skipped = [r for r in records if r["action"] == "skip"]
    in_scope = inserts + exists + [r for r in skipped if r["kind"]]

    rows = [["raw_id", "action", "emp_code", "kind", "name", "email", "designation",
             "department", "sheet_colour", "date_of_joining", "sheet_date_of_leaving",
             "contract_ends", "date_of_exit", "user_status", "l1_manager", "l2_manager",
             "reason/note"]]
    for r in sorted(in_scope, key=lambda x: (x["action"], x["emp_code"] or x["raw_id"])):
        rows.append([
            r["raw_id"], r["action"], r["emp_code"], r["kind"] or "",
            f"{r['first']} {r['last']}".strip(), r["email"], r["designation"],
            r["department"],
            "" if r.get("red") is None else ("red" if r["red"] else "normal"),
            r["doj"] or "", r["dol"] or "",
            r.get("ends_on") or "", "" if r.get("active") else (r["dol"] or ""),
            r.get("user_status", ""), r["mgr1"], r["mgr2"],
            r.get("reason") or r.get("note", ""),
        ])

    n_active = sum(1 for r in inserts if r["active"])
    lines = [
        "# Intern / Contractual import report",
        "",
        f"- Source: `{source}`",
        f"- Mode: **{'COMMIT' if committed else 'DRY RUN'}** ({mode} accounts)",
        f"- Active/exit rule: **{'red row = exited (sheet colour coding)' if colour_rule else f'contract end date vs {as_of}'}**",
        f"- Rows in the sheet: {len(records)}",
        f"- Intern/contract rows found: {len(in_scope)}",
        f"- **Inserted: {len(inserts)}** "
        f"(contract {sum(1 for r in inserts if r['kind'] == 'contract')}, "
        f"intern {sum(1 for r in inserts if r['kind'] == 'intern')}; "
        f"active {n_active}, exit {len(inserts) - n_active})",
        f"- Already in the DB (untouched): {len(exists)}",
        f"- Skipped in-scope rows: {len([r for r in skipped if r['kind']])}",
        f"- Out-of-scope rows (permanent/other, never touched): {len([r for r in skipped if not r['kind']])}",
        "",
        "## Inserted",
        "",
        "| emp_code | name | kind | designation | joined | contract ends | exited | account |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for r in sorted(inserts, key=lambda x: x["emp_code"]):
        lines.append(f"| `{r['emp_code']}` | {r['first']} {r['last']} | {r['kind']} | "
                     f"{r['designation']} | {r['doj'] or ''} | {r.get('ends_on') or ''} | "
                     f"{'' if r['active'] else (r['dol'] or '')} | {r['user_status']} |")
    odd = [r for r in in_scope if r.get("disagrees")]
    lines += ["", "## Colour vs contract end date — worth an eyeball", ""]
    lines += ["_The colour won: these rows are imported as stated in the 'as imported' column._", ""] if odd else []
    lines += [f"- `{r['emp_code'] or r['raw_id']}` {r['first']} {r['last']} — sheet is "
              f"**{'red' if r['red'] else 'normal'}**, DateOfLeaving "
              f"{r['dol'] or 'empty'} → as imported: **{r['user_status']}**"
              for r in odd] or ["_None — colour and dates agree._"]

    rehires = [r for r in inserts if r.get("rehire_of") or r.get("_link_uid")]
    lines += ["", "## Rehire stints — employee record kept, linked to one existing login", ""]
    lines += [f"- `{r['emp_code']}` {r['first']} {r['last']} <{r['email']}> — {r['note']}"
              for r in rehires] or ["_None._"]
    lines += ["", "## Already in the DB — not inserted, nothing updated", ""]
    lines += [f"- `{r['raw_id']}` {r['first']} {r['last']} <{r['email']}> — {r['reason']}"
              for r in exists] or ["_None._"]
    lines += ["", "## Skipped intern/contract rows", ""]
    lines += [f"- `{r['raw_id']}` <{r['email']}> — {r['reason']}"
              for r in skipped if r["kind"]] or ["_None._"]
    lines += ["", "## Unresolved manager names (left null)", ""]
    lines += [f"- {n}" for n in sorted(unresolved)] or ["_None._"]
    lines += ["", f"_Per-row detail: {REPORT_CSV.name}_", ""]

    with open(REPORT_CSV, "w", encoding="utf-8", newline="") as f:
        csv.writer(f).writerows(rows)
    REPORT_MD.write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines[:14]))
    print(f"\nReport written:\n  {REPORT_MD}\n  {REPORT_CSV}")


# ── Entry ─────────────────────────────────────────────────────────────────
async def main():
    ap = argparse.ArgumentParser(description="Import only intern + contractual employees")
    ap.add_argument("--commit", action="store_true", help="write to the DB (default: dry run)")
    ap.add_argument("--revert", action="store_true", help="delete everything THIS importer created, then exit")
    ap.add_argument("--pending", action="store_true",
                    help="create active employees INACTIVE (no password) instead of ACTIVE/test@123")
    ap.add_argument("--activation", action="store_true",
                    help="create active employees INACTIVE and enqueue activation emails")
    ap.add_argument("--emit-events", action="store_true",
                    help="publish department/employee.created events to the outbox")
    ap.add_argument("--skip-existing-email", action="store_true",
                    help="do NOT import a stint whose email already has a user; report it instead "
                         "(default: insert the employee record linked to that login as a rehire)")
    ap.add_argument("--file", help="source sheet (.xlsx or .csv); overrides --data-dir")
    ap.add_argument("--data-dir", help="folder holding the sheet (picked by name)")
    ap.add_argument("--org-id", help="organisation id (default: resolve from the org admin email)")
    ap.add_argument("--prefix", default=DEFAULT_PREFIX, help=f"emp_code prefix (default {DEFAULT_PREFIX})")
    ap.add_argument("--as-of", help="YYYY-MM-DD for the date fallback when the source has no colours "
                                    "(default: today)")
    ap.add_argument("--ignore-colours", action="store_true",
                    help="ignore the red/not-red colour coding and decide active/exit from the "
                         "contract end date instead")
    args = ap.parse_args()
    mode = "activation" if args.activation else ("pending" if args.pending else "active")
    as_of = parse_date(args.as_of) or date.today()
    if args.as_of and not parse_date(args.as_of):
        print(f"ERROR: could not read --as-of '{args.as_of}' (use YYYY-MM-DD)")
        sys.exit(1)

    source = None
    records: list[dict] = []
    maxes: dict = {"C": 0, "I": 0}
    if not args.revert:
        if args.file:
            source = Path(args.file)
        elif args.data_dir:
            source = find_file(args.data_dir, ["import-template", "employee"])
        else:
            source = DEFAULT_FILE
        if not source.exists():
            print(f"ERROR: source file not found at {source}")
            sys.exit(1)
        rows, has_colours = read_rows_colored(source)
        use_colour = has_colours and not args.ignore_colours
        records, maxes = parse_rows(rows, args.prefix, as_of, use_colour=use_colour)
        n_scope = len([r for r in records if r["kind"]])
        n_active = len([r for r in records if r["kind"] and r.get("active")])
        rule = ("red = exited (sheet colour coding)" if use_colour
                else f"contract end date vs {as_of} (no colours in the source)")
        print(f"Read {len(rows)} rows from {source.name}; {n_scope} are intern/contract "
              f"({n_active} active, {n_scope - n_active} exited) — rule: {rule}.")

    if not settings.MONGODB_URL:
        print("ERROR: MONGODB_URL not configured.")
        sys.exit(1)
    client = AsyncIOMotorClient(settings.MONGODB_URL)
    await init_beanie(
        database=client.get_default_database(),
        document_models=[UserDocument, OrganisationDocument, BusinessUnitDocument,
                         DepartmentDocument, DesignationDocument, EmployeeDocument,
                         AddressDocument, MasterDataDocument, OutboxEventDocument],
    )

    if args.org_id:
        org_id = PydanticObjectId(args.org_id)
    else:
        admin = await UserDocument.get_motor_collection().find_one(
            {"email": ORG_ADMIN_EMAIL.lower()}, {"organisation_id": 1})
        org_id = PydanticObjectId(admin["organisation_id"]) if admin else PydanticObjectId(ORG_ID_FALLBACK)
    org = await OrganisationDocument.get(ObjectId(org_id))
    if not org:
        print(f"ERROR: org not found (org_id={org_id}).")
        client.close()
        sys.exit(1)
    print(f"Org: {org.legal_name} (id={org_id})")

    if args.revert:
        await cleanup(org_id)
        print("Revert complete — every doc created by this importer removed.")
        client.close()
        return

    bu = await pick_bu(org_id, args.prefix)
    await mark_existing(records, org_id, args.prefix,
                        skip_existing_email=args.skip_existing_email)

    if args.commit:
        unresolved = await run(records, maxes, org_id, bu, mode=mode, emit_events=args.emit_events)
    else:
        users = [u async for u in UserDocument.get_motor_collection().find(
            {"organisation_id": ObjectId(org_id)}, {"first_name": 1, "last_name": 1})]
        names = [f"{u.get('first_name') or ''} {u.get('last_name') or ''}" for u in users]
        names += [f"{r['first']} {r['last']}" for r in records if r["action"] == "insert"]
        known = {norm_name(n) for n in names}
        skeys = {sorted_key(n) for n in names}
        unresolved = dryrun_managers(records, known, skeys)
        print("DRY RUN — no DB writes.")

    write_report(records, unresolved, committed=args.commit, source=source, mode=mode,
                 as_of=as_of, colour_rule=use_colour)
    client.close()


if __name__ == "__main__":
    asyncio.run(main())
