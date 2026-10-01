"""Populate the ``payslips`` collection from an uploaded salary file (xlsx or csv).

On upload: read the grid (the signature sheet for a workbook, or the CSV rows),
map columns **by header name** (so they can be reshuffled), resolve each row's
``emp_code`` to the IAM ``user_id``, then upsert
one ``payslips`` document per employee. Every amount column (current + all
``* YTD`` + CTC/Per Month/Fixed Basic/LOP/net) is captured at its exact value
into an encrypted ``extra`` snapshot; a display subset also lands in
``earnings``/``deductions``. Statutory/bank ids and amounts are encrypted at rest.

Resilience: if IAM / the broker is unreachable the file stays stored and the
result is reported as ``deferred`` (re-runnable later).
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime

import openpyxl
from beanie import PydanticObjectId
from pymongo.errors import DuplicateKeyError

from src.correlation import audit_create, get_correlation_id, stamp_modified
from src.logger import logger
from src.payslips import constants
from src.payslips.models import PayslipDocument, UploadedPayslipDocument
from src.payslips.schemas import (
    Deductions,
    Earnings,
    PayslipCreate,
    PayslipImportResult,
    PayslipRowError,
    PayslipRowValidation,
    PayslipValidationResult,
)
from src.rabbitmq.iam_rpc import IamUnavailable, resolve_employees

__all__ = ["populate_from_upload", "validate_workbook"]

# Validation stays responsive even if IAM is slow — bound its employee lookup.
_VALIDATE_RPC_TIMEOUT = 15.0

_UNMAPPED = object()  # sentinel: header not in the known map (capture raw)

# Normalised sheet header (lowercase, single-spaced) -> canonical field.
# ``None`` = intentionally ignored. Unknown headers are captured raw into ``extra``.
_HEADER_MAP: dict[str, str | None] = {
    "s.no": None,
    "sil id": "emp_code",
    "name of the employee": "full_name",
    "designation": "designation",
    "ctc": "ctc",
    "per month": "per_month",
    "fixed basic salary": "fixed_basic",
    "basic": "basic",
    "basic ytd": "basic_ytd",
    "variable pay / performance incentive": "variable_pay",
    # Legacy (<= May 2026) earning columns, retired in the June 2026 template.
    "retention incentive": "retention_incentive",
    "uniform allowance": "uniform_allowance",
    "telephone & internet allowance.": "telephone_or_mobile",
    "professional pursuit allowance": "magazines",
    "days in month": "standard_days",
    "no of days worked": "days_worked",
    "lop": "lop",
    "salary earned": "salary_earned",
    "hra": "hra",
    "hra ytd": "hra_ytd",
    "arrears": "arrears",
    "arrears ytd": "arrears_ytd",
    "lta": "lta",
    "lta ytd": "lta_ytd",
    "incentive /projects allowance": "incentive_projects_allowance",
    "incentive /projects allowance ytd": "incentive_projects_allowance_ytd",
    "totalearnings": "total_earnings",
    "income tax": "income_tax",
    "income tax ytd": "income_tax_ytd",
    "pt": "professional_tax",
    "pt ytd": "professional_tax_ytd",
    "pf": "provident_fund",
    "pf ytd": "provident_fund_ytd",
    "esi": "esi",  # legacy (<= May 2026) deduction column
    "health insurance premium 25%": "health_insurance_premium",
    "health insurance premium 25% ytd": "health_insurance_premium_ytd",
    "gmc premium parental": "gmc_premium",
    "gmc premium parental ytd": "gmc_premium_ytd",
    "othersdeductions": "other_deductions",
    "advance to staff": "salary_advance",
    "advance to staff ytd": "salary_advance_ytd",
    "totaldeductions": "total_deductions",
    "net salary": "net_salary",
    "pf no": "pf_no",
    "pan no": "pan_no",
    "bank name": "bank_name",
    "account no": "account_no",
    "date of joining": "date_of_joining",
    "uan no": "uan_number",
    "gender": "gender",
    # the trailing columns to ignore (incl. the legacy typo'd headers)
    "pf employer share": None,
    "gratuity": None,
    "gratuty": None,
    "variable pay": None,
    "varibale salary": None,
}

# Canonical fields that are clear text / counts, not encrypted amounts.
_TEXT_FIELDS = {"emp_code", "full_name", "designation", "date_of_joining", "gender"}
_STATUTORY_FIELDS = {"pf_no", "pan_no", "bank_name", "account_no", "uan_number"}
_DAY_FIELDS = {"standard_days", "days_worked"}
# Legacy-only amount columns (<= May 2026), absent from the current template.
_LEGACY_FIELDS = {"retention_incentive", "uniform_allowance", "telephone_or_mobile", "magazines", "esi"}
# Current-only columns (June 2026 onward): the YTD series + variable pay. Absent
# from the legacy template.
_MAPPED_FIELDS = {v for v in _HEADER_MAP.values() if v}
_NEW_ONLY_FIELDS = {f for f in _MAPPED_FIELDS if f.endswith("_ytd")} | {"variable_pay"}
# Columns we expect for each template — used only to report ``missing_columns``
# (a missing column never invalidates a file). The legacy set keeps its own
# columns and drops the YTD/variable ones; the current set is the reverse.
_EXPECTED_FIELDS_NEW = _MAPPED_FIELDS - _LEGACY_FIELDS
_EXPECTED_FIELDS_OLD = _MAPPED_FIELDS - _NEW_ONLY_FIELDS
_EXPECTED_FIELDS = _EXPECTED_FIELDS_NEW  # default when no period is given
# Everything mapped that isn't text/statutory/count is an amount -> float -> the
# encrypted ``extra`` snapshot (both templates' amount columns, so all project).
_AMOUNT_FIELDS = _MAPPED_FIELDS - _TEXT_FIELDS - _STATUTORY_FIELDS - _DAY_FIELDS

# Display ``Earnings``/``Deductions`` projected from the amount snapshot.
_EARNINGS_FROM = {
    "basic_salary": "salary_earned",
    "hra": "hra",
    "LTA": "lta",
    "arrears": "arrears",
    "incentive_or_project_allowwance": "incentive_projects_allowance",
    # Legacy (<= May 2026) earning lines; 0 for newer periods that lack the column.
    "uniform_allowance": "uniform_allowance",
    "telephone_or_mobile": "telephone_or_mobile",
    "magazines": "magazines",
    "retention_incentive": "retention_incentive",
    "total": "total_earnings",
}
_DEDUCTIONS_FROM = {
    "income_tax": "income_tax",
    "professional_tax": "professional_tax",
    "provident_fund": "provident_fund",
    "health_insurance_premium": "health_insurance_premium",
    "gmc_premium": "gmc_premium",
    "other_deductions": "other_deductions",
    "salary_advance": "salary_advance",
    "esi": "esi",  # legacy (<= May 2026) deduction line
    "total": "total_deductions",
}


# ── Cell coercion ────────────────────────────────────────────────────────────


def _normalize(header: object) -> str:
    return " ".join(str(header).strip().lower().split()) if header is not None else ""


def _nonblank(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def _to_float(value: object) -> float:
    """Exact numeric value (no rounding). Blank -> 0.0; raises ValueError if bad."""
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    return float(str(value).strip().replace(",", ""))


def _to_int_or_none(value: object) -> int | None:
    if not _nonblank(value):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return int(round(float(str(value).strip().replace(",", ""))))


def _to_text(value: object) -> str | None:
    """A clean string for ids/names/dates (no ``.0`` on whole numbers)."""
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text or None


def _first(*values: object) -> str | None:
    for value in values:
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


# ── Parsing ──────────────────────────────────────────────────────────────────


def _looks_like_zip(body: bytes) -> bool:
    """xlsx/xlsm are ZIP archives (PK magic); anything else is treated as CSV/text."""
    return body[:4] in (b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")


def _signature_index(rows: list) -> int | None:
    """1-based index of the header row (matching the signature) within the first 8 rows."""
    for r_idx, row in enumerate(rows[:8], start=1):
        norms = {_normalize(c) for c in row if _nonblank(c)}
        if all(sig in norms for sig in constants.SHEET_SIGNATURE):
            return r_idx
    return None


def _read_grid(body: bytes) -> tuple[list[list], str | None]:
    """Read the upload into a grid of rows (each a list of cell values).

    Handles both xlsx/xlsm workbooks and CSV. For a multi-sheet workbook the sheet
    carrying the signature header is chosen. CSV is decoded UTF-8 (BOM-tolerant),
    falling back to latin-1. Returns ``(grid, file_error)``.
    """
    if _looks_like_zip(body):
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
        except Exception as exc:  # corrupt / not a real xlsx
            return [], f"Could not read the .xlsx file: {exc}"
        sheets = workbook.worksheets
        for sheet in sheets:
            grid = [list(row) for row in sheet.iter_rows(values_only=True)]
            if _signature_index(grid) is not None:
                return grid, None
        # No sheet matched the signature — hand back the first so the caller reports it.
        return ([list(row) for row in sheets[0].iter_rows(values_only=True)] if sheets else []), None

    try:
        text = body.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = body.decode("latin-1", errors="replace")
    return [list(row) for row in csv.reader(io.StringIO(text))], None


def _earnings_from(extra: dict) -> Earnings:
    return Earnings(**{field: extra.get(src, 0.0) for field, src in _EARNINGS_FROM.items()})


def _deductions_from(extra: dict) -> Deductions:
    return Deductions(**{field: extra.get(src, 0.0) for field, src in _DEDUCTIONS_FROM.items()})


def _parse_workbook(body: bytes, expected_fields: set[str] | None = None) -> tuple[list[dict], str | None, list[str]]:
    """Parse an uploaded xlsx/xlsm workbook or CSV into per-row dicts.

    Returns ``(rows, file_level_error, missing_columns)``. ``missing_columns`` is
    reported against ``expected_fields`` (defaults to the current template) — it
    is informational only; absent amount columns default to 0.
    """
    expected = expected_fields if expected_fields is not None else _EXPECTED_FIELDS
    grid, file_error = _read_grid(body)
    if file_error:
        return [], file_error, []

    header_row = _signature_index(grid)
    if header_row is None:
        return [], "No payslip data sheet found (needs the 'SIL ID' and 'NET SALARY' columns).", []
    headers = grid[header_row - 1]

    col_canonical: dict[int, str] = {}
    col_unknown: dict[int, str] = {}
    for idx, raw_header in enumerate(headers):
        norm = _normalize(raw_header)
        if not norm:
            continue
        mapped = _HEADER_MAP.get(norm, _UNMAPPED)
        if mapped is None:
            continue
        if mapped is _UNMAPPED:
            col_unknown[idx] = norm
        else:
            col_canonical[idx] = mapped
    present = set(col_canonical.values())
    missing_columns = sorted(expected - present)

    rows: list[dict] = []
    for line_no, raw in enumerate(grid[header_row:], start=header_row + 1):
        record = {col_canonical[i]: raw[i] for i in col_canonical if i < len(raw)}
        unknown = {col_unknown[i]: raw[i] for i in col_unknown if i < len(raw)}
        if not any(_nonblank(v) for v in (*record.values(), *unknown.values())):
            continue  # fully blank line
        emp_code = _to_text(record.get("emp_code"))
        if not emp_code:
            rows.append({"row_num": line_no, "emp_code": None, "error": "Missing employee id (SIL ID)."})
            continue
        try:
            extra = {field: _to_float(record.get(field)) for field in _AMOUNT_FIELDS}
            for key, value in unknown.items():  # keep any unexpected column too
                extra[key] = _to_float(value) if isinstance(value, (int, float)) else _to_text(value)
            earnings = _earnings_from(extra)
            deductions = _deductions_from(extra)
        except (ValueError, TypeError) as exc:
            rows.append({"row_num": line_no, "emp_code": emp_code, "error": f"Non-numeric amount: {exc}"})
            continue
        rows.append(
            {
                "row_num": line_no,
                "emp_code": emp_code,
                "earnings": earnings,
                "deductions": deductions,
                "extra": extra,
                "full_name": _to_text(record.get("full_name")),
                "designation": _to_text(record.get("designation")),
                "date_of_joining": _to_text(record.get("date_of_joining")),
                "gender": _to_text(record.get("gender")),
                "standard_days": _to_int_or_none(record.get("standard_days")),
                "days_worked": _to_int_or_none(record.get("days_worked")),
                "statutory": {f: _to_text(record.get(f)) for f in _STATUTORY_FIELDS},
            }
        )
    return rows, None, missing_columns


def _build_create(row: dict, info: dict, upload: UploadedPayslipDocument) -> PayslipCreate:
    """Merge a parsed row with IAM info (the sheet wins; IAM fills gaps)."""
    stat = row["statutory"]
    return PayslipCreate(
        payslip_file_id=str(upload.id),
        user_id=str(info["user_id"]),
        emp_code=row["emp_code"],
        month=upload.month,
        year=upload.year,
        full_name=_first(row.get("full_name"), info.get("full_name")),
        designation=_first(row.get("designation"), info.get("designation")),
        date_of_joining=_first(row.get("date_of_joining"), info.get("date_of_joining")),
        gender=_first(row.get("gender"), info.get("gender")),
        earnings=row["earnings"],
        deductions=row["deductions"],
        extra=row["extra"],
        standard_days=row["standard_days"],
        days_worked=row["days_worked"],
        uan_number=_first(stat["uan_number"], info.get("uan_number")),
        pf_no=_first(stat["pf_no"], info.get("pf_no")),
        pan_no=_first(stat["pan_no"], info.get("pan_no")),
        bank_name=_first(stat["bank_name"], info.get("bank_name")),
        account_no=_first(stat["account_no"], info.get("account_no")),
    )


# ── Persistence ──────────────────────────────────────────────────────────────


async def _upsert(data: PayslipCreate, organisation_id: PydanticObjectId, business_unit_id: PydanticObjectId) -> str:
    """Insert a new payslip, or refresh the existing one for (org, user, period).

    Matched on the unique ``(user_id, year, month)`` within the org — *not* the
    business unit — so a re-upload matches regardless of which BU stamped it.
    ``business_unit_id`` is the uploader's (so uploads appear in their BU-scoped
    list), even when the employee belongs to another BU in the same org.
    """
    existing = await PayslipDocument.find_one(
        PayslipDocument.organisation_id == organisation_id,
        PayslipDocument.user_id == PydanticObjectId(data.user_id),
        PayslipDocument.year == data.year,
        PayslipDocument.month == data.month,
    )
    if existing is not None:
        existing.apply_payload(data)
        stamp_modified(existing)
        await existing.save()
        return "updated"
    doc = PayslipDocument.from_create(
        data,
        organisation_id=organisation_id,
        business_unit_id=business_unit_id,
        correlation_id=get_correlation_id(),
        **audit_create(),
    )
    await doc.insert()
    return "created"


def _org_mismatch(info: dict, organisation_id: PydanticObjectId) -> bool:
    """True if the resolved employee is outside the caller's organisation.

    Business unit is intentionally not checked — employees are accepted org-wide.
    """
    org = info.get("organisation_id")
    return org is not None and str(org) != str(organisation_id)


async def populate_from_upload(
    upload: UploadedPayslipDocument,
    body: bytes,
    organisation_id: PydanticObjectId,
    business_unit_id: PydanticObjectId,
    *,
    resolve=resolve_employees,
) -> PayslipImportResult:
    """Parse ``body`` and upsert the resulting payslips for ``upload``'s period."""
    parsed, file_error, _ = _parse_workbook(body)
    if file_error:
        return PayslipImportResult(status="skipped", message=file_error)
    if not parsed:
        return PayslipImportResult(status="completed", total_rows=0, message="No data rows found.")

    errors = [
        PayslipRowError(row_num=r["row_num"], emp_code=r.get("emp_code"), message=r["error"])
        for r in parsed
        if r.get("error")
    ]
    valid = [r for r in parsed if not r.get("error")]
    emp_codes = sorted({r["emp_code"] for r in valid})

    try:
        # Empty BU -> IAM searches the whole organisation (employees accepted org-wide).
        resolved = await resolve(emp_codes, str(organisation_id), "")
    except IamUnavailable as exc:
        logger.warning("payslip_import.iam_unavailable", upload_id=str(upload.id), error=str(exc))
        return PayslipImportResult(
            status="deferred",
            total_rows=len(parsed),
            message=(
                f"Employee lookup is unavailable ({exc}). The file is stored, but payslip "
                "rows were not populated — retry once IAM is reachable."
            ),
        )

    created = updated = skipped = 0
    for row in valid:
        emp_code = row["emp_code"]
        info = resolved.get(emp_code)
        if not info or not info.get("user_id"):
            skipped += 1
            errors.append(
                PayslipRowError(
                    row_num=row["row_num"],
                    emp_code=emp_code,
                    message="Employee not found in IAM for your organisation.",
                )
            )
            continue
        if _org_mismatch(info, organisation_id):
            skipped += 1
            errors.append(
                PayslipRowError(
                    row_num=row["row_num"],
                    emp_code=emp_code,
                    message="Employee belongs to a different organisation.",
                )
            )
            continue
        try:
            outcome = await _upsert(_build_create(row, info, upload), organisation_id, business_unit_id)
        except DuplicateKeyError:
            skipped += 1
            errors.append(
                PayslipRowError(row_num=row["row_num"], emp_code=emp_code, message="Write conflict; please retry.")
            )
            continue
        except Exception as exc:
            logger.error("payslip_import.row_failed", row=row["row_num"], emp_code=emp_code, error=repr(exc))
            skipped += 1
            errors.append(
                PayslipRowError(row_num=row["row_num"], emp_code=emp_code, message=f"Could not save row: {exc}")
            )
            continue
        created += outcome == "created"
        updated += outcome == "updated"

    logger.info(
        "payslip_import.done",
        upload_id=str(upload.id),
        total=len(parsed),
        created=created,
        updated=updated,
        skipped=skipped,
    )
    return PayslipImportResult(
        status="completed",
        total_rows=len(parsed),
        created=created,
        updated=updated,
        skipped=skipped,
        errors=errors,
    )


async def validate_workbook(
    body: bytes,
    organisation_id: PydanticObjectId,
    business_unit_id: PydanticObjectId,
    *,
    old_template: bool = False,
    resolve=resolve_employees,
) -> PayslipValidationResult:
    """Dry-run validate a salary workbook — per-row status/issues, nothing stored.

    A missing template column is never an error: it's reported under
    ``missing_columns`` and those amounts default to 0. ``old_template`` reports
    ``missing_columns`` against the legacy (<= May 2026) column set so a legacy
    file isn't flagged for the YTD columns it never had. Employee existence is
    checked org-wide against IAM best-effort.
    """
    expected = _EXPECTED_FIELDS_OLD if old_template else _EXPECTED_FIELDS_NEW
    parsed, file_error, missing_columns = _parse_workbook(body, expected)
    if file_error:
        return PayslipValidationResult(valid=False, file_errors=[file_error])

    counts: dict[str, int] = {}
    for row in parsed:
        code = row.get("emp_code")
        if code and not row.get("error"):
            counts[code] = counts.get(code, 0) + 1

    resolved: dict[str, dict | None] = {}
    employees_checked = True
    try:
        resolved = await resolve(sorted(counts), str(organisation_id), "", timeout=_VALIDATE_RPC_TIMEOUT)
    except IamUnavailable as exc:
        employees_checked = False
        logger.warning("payslip_validate.iam_unavailable", error=str(exc))

    rows: list[PayslipRowValidation] = []
    for row in parsed:
        code = row.get("emp_code")
        issues: list[str] = []
        if row.get("error"):
            issues.append(row["error"])
        else:
            if counts.get(code, 0) > 1:
                issues.append("Duplicate emp_code in the file; the last row would overwrite the earlier ones.")
            if employees_checked:
                info = resolved.get(code)
                if not info or not info.get("user_id"):
                    issues.append("Employee not found in IAM for your organisation.")
                elif _org_mismatch(info, organisation_id):
                    issues.append("Employee belongs to a different organisation.")
        rows.append(
            PayslipRowValidation(
                row_num=row["row_num"],
                emp_code=code,
                status="error" if issues else "valid",
                issues=issues,
            )
        )

    error_rows = sum(1 for row in rows if row.status == "error")
    return PayslipValidationResult(
        valid=True,
        total_rows=len(rows),
        valid_rows=len(rows) - error_rows,
        error_rows=error_rows,
        missing_columns=missing_columns,
        employees_checked=employees_checked,
        rows=rows,
    )
