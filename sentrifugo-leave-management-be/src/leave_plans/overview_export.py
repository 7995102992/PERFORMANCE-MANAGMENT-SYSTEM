import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ─── Style palette ───────────────────────────────────────────────────────────

_TITLE_FONT = Font(bold=True, size=14, color="1F3864")
_SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
_SECTION_FILL = PatternFill("solid", fgColor="366092")
_HEADER_FONT = Font(bold=True, size=10, color="1F3864")
_HEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BU_FONT = Font(bold=True, size=10, color="1F3864")
_BU_FILL = PatternFill("solid", fgColor="E2EFDA")
_NUM_ALIGN = Alignment(horizontal="right", vertical="center")
_LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
_THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
)
_TOTAL_FONT = Font(bold=True, size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December",
]


def _format_label(s: str) -> str:
    return s.replace("_", " ").replace("-", " ").title()


def _set_row(ws, row: int, values: list, *, font=None, fill=None, alignment=None, border=None):
    for col_idx, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if alignment:
            cell.alignment = alignment
        elif isinstance(val, (int, float)):
            cell.alignment = _NUM_ALIGN
        else:
            cell.alignment = _LEFT_ALIGN
        if border:
            cell.border = border


def _section_banner(ws, row: int, text: str, col_count: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    return row + 1


def generate_overview_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leave Plan Overview"

    status_keys = sorted({
        s
        for bu in data.get("business_units", [])
        for dept in bu.get("departments", [])
        for s in dept.get("employee_counts", {}).get("by_status", {}).keys()
    })

    # Column layout: BU | Department | Total Employees | <status1> | <status2> | ... | Total Days Allocated | <status1> days | ...
    status_labels = [_format_label(s) for s in status_keys]
    emp_headers = ["Total Employees"] + status_labels
    alloc_headers = ["Total Days Allocated"] + [f"{l} Days" for l in status_labels]
    all_headers = ["Business Unit", "Department"] + emp_headers + alloc_headers
    col_count = len(all_headers)

    row = 1

    # ── Plan title ───────────────────────────────────────────────────────
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=data.get("plan_name", "Leave Plan"))
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28
    row += 1

    # ── Plan details ─────────────────────────────────────────────────────
    cal_month = data.get("calendar_start_month")
    cal_label = MONTH_NAMES[cal_month - 1] if cal_month and 1 <= cal_month <= 12 else "—"

    details = [
        ("Plan Year Starts", cal_label),
    ]

    prob = data.get("probation_config")
    if prob and prob.get("enabled"):
        cm = _format_label(prob.get("credit_mode") or "—")
        dur = prob.get("probation_duration_months")
        label = cm + (f", {dur} months" if dur else "")
        details.append(("Probation", label))

    np_cfg = data.get("notice_period_config")
    if np_cfg and np_cfg.get("mode"):
        details.append(("Notice Period", _format_label(np_cfg["mode"])))

    leave_types = data.get("leave_types", [])

    for label, value in details:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=row, column=1).alignment = _LEFT_ALIGN
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=col_count)
        ws.cell(row=row, column=2, value=value).alignment = _LEFT_ALIGN
        row += 1

    row += 1  # blank row

    # ── Assignment breakdown ─────────────────────────────────────────────
    row = _section_banner(ws, row, "Assignment Breakdown", col_count)

    # Header row
    ws.row_dimensions[row].height = 26
    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    row += 1

    # Data rows grouped by BU → Dept
    for bu in data.get("business_units", []):
        bu_name = bu.get("name", "—")
        bu_emp = bu.get("employee_counts", {})
        bu_alloc = bu.get("leaves_allocated_by_employee_type", {})

        for dept_idx, dept in enumerate(bu.get("departments", [])):
            dept_emp = dept.get("employee_counts", {})
            dept_alloc = dept.get("leaves_allocated_by_employee_type", {})
            by_status = dept_emp.get("by_status", {})

            values = [
                bu_name if dept_idx == 0 else "",
                dept.get("name", "—"),
                dept_emp.get("total", 0),
            ]
            for sk in status_keys:
                values.append(by_status.get(sk, 0))
            values.append(dept.get("total_leaves_allocated", 0))
            for sk in status_keys:
                values.append(dept_alloc.get(sk, 0))

            for col_idx, val in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                if col_idx <= 2:
                    cell.alignment = _LEFT_ALIGN
                else:
                    cell.alignment = _NUM_ALIGN
                if dept_idx == 0 and col_idx == 1:
                    cell.font = _BU_FONT
                    cell.fill = _BU_FILL
            row += 1

        # BU subtotal
        bu_vals = [
            "",
            f"{bu_name} — Total",
            bu_emp.get("total", 0),
        ]
        for sk in status_keys:
            bu_vals.append(bu_emp.get("by_status", {}).get(sk, 0))
        bu_vals.append(bu.get("total_leaves_allocated", 0))
        for sk in status_keys:
            bu_vals.append(bu_alloc.get(sk, 0))

        for col_idx, val in enumerate(bu_vals, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = _TOTAL_FONT
            cell.fill = _TOTAL_FILL
            cell.border = _THIN_BORDER
            cell.alignment = _NUM_ALIGN if col_idx > 2 else _LEFT_ALIGN
        row += 1

    # Grand total
    totals = data.get("totals", {})
    alloc_by_type = totals.get("leaves_allocated_by_employee_type", {})
    grand_vals = [
        "",
        "Grand Total",
        totals.get("total_employees", 0),
    ]
    for sk in status_keys:
        grand_vals.append(0)
    grand_vals.append(totals.get("total_leaves_allocated", 0))
    for sk in status_keys:
        grand_vals.append(alloc_by_type.get(sk, 0))

    # Sum per-status emp counts from BU data for the grand row
    for bu in data.get("business_units", []):
        by_status = bu.get("employee_counts", {}).get("by_status", {})
        for i, sk in enumerate(status_keys):
            grand_vals[3 + i] += by_status.get(sk, 0)

    ws.row_dimensions[row].height = 24
    for col_idx, val in enumerate(grand_vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if col_idx > 2 else _LEFT_ALIGN
    row += 2

    # ── Employees by Employment Type ─────────────────────────────────────
    est_counts = data.get("employment_status_counts") or []
    if est_counts:
        row = _section_banner(ws, row, "Employees by Employment Type", col_count)
        for col_idx, header in enumerate(["Employment Type", "Employees"], start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _LEFT_ALIGN if col_idx == 1 else _NUM_ALIGN
            cell.border = _THIN_BORDER
        row += 1
        for item in est_counts:
            c1 = ws.cell(row=row, column=1, value=item.get("label") or item.get("key", "—"))
            c1.alignment = _LEFT_ALIGN
            c1.border = _THIN_BORDER
            c2 = ws.cell(row=row, column=2, value=item.get("count", 0))
            c2.alignment = _NUM_ALIGN
            c2.border = _THIN_BORDER
            row += 1
        row += 1

    # ── Leave Types table ────────────────────────────────────────────────
    # Each leave type now owns its allocation, accrual, carry-forward and
    # encashment — so the export reflects that directly.
    if leave_types:
        row = _section_banner(ws, row, "Leave Types", col_count)
        lt_headers = ["#", "Leave Type", "Allocated / Year", "Accrual", "Carry Forward", "Encashable"]
        for col_idx, header in enumerate(lt_headers, start=1):
            cell = ws.cell(row=row, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _CENTER_ALIGN if col_idx == 1 else _LEFT_ALIGN
            cell.border = _THIN_BORDER
        row += 1
        for idx, lt in enumerate(leave_types, start=1):
            unit = str(lt.get("unit", "DAYS")).lower()
            alloc = lt.get("annual_allocated")
            alloc_str = f"{alloc} {unit}" if alloc is not None else "—"

            if lt.get("is_statutory"):
                accrual_str, carry_str, encash_str = "Annual (statutory)", "No", "No"
            elif not (lt.get("is_paid", True) and lt.get("is_paid_leave", True)):
                # Unpaid when EITHER paid flag is off (the two are redundant
                # duplicates) — keeps this export consistent with the overview
                # service and the LOP decision, which both treat either-off as unpaid.
                accrual_str, carry_str, encash_str = "—", "No", "No"
            else:
                freq = lt.get("accrual_frequency")
                accrual_str = _format_label(freq) if freq else "—"
                if lt.get("carry_forward"):
                    cnt = lt.get("carry_forward_count")
                    carry_str = f"Up to {cnt} {unit}" if cnt is not None else "Yes"
                else:
                    carry_str = "No"
                if lt.get("encashable"):
                    pct = lt.get("encash_percentage")
                    encash_str = f"{pct}% encash" if pct is not None else "Yes"
                else:
                    encash_str = "No"

            for col_idx, val in enumerate(
                [idx, lt.get("name", "—"), alloc_str, accrual_str, carry_str, encash_str], start=1
            ):
                cell = ws.cell(row=row, column=col_idx, value=val)
                cell.border = _THIN_BORDER
                cell.alignment = _CENTER_ALIGN if col_idx == 1 else _LEFT_ALIGN
            row += 1

    # ── Column widths ────────────────────────────────────────────────────
    widths = [22, 26] + [16] * (len(emp_headers)) + [20] + [16] * len(status_keys)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "C1"

    _add_employee_breakdown_sheet(wb, data)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _add_employee_breakdown_sheet(wb: Workbook, data: dict) -> None:
    employees = data.get("employees") or []
    leave_types = data.get("leave_types") or []
    if not employees:
        return

    ws = wb.create_sheet(title="Employee Breakdown")

    lt_names = [lt["name"] for lt in leave_types]
    base_headers = [
        "#",
        "Emp Code",
        "Name",
        "Email",
        "Business Unit",
        "Department",
        "Status",
        "Date of Joining",
        "Total Days Allocated",
    ]
    all_headers = base_headers + lt_names
    col_count = len(all_headers)

    row = 1

    # ── Title ────────────────────────────────────────────────────────────
    title = f"{data.get('plan_name', 'Leave Plan')} — Employee Breakdown"
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28
    row += 2  # blank row after title

    # ── Section banner ───────────────────────────────────────────────────
    row = _section_banner(ws, row, "Employees", col_count)

    # ── Header row ───────────────────────────────────────────────────────
    ws.row_dimensions[row].height = 26
    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    header_row = row
    row += 1

    # ── Data rows ────────────────────────────────────────────────────────
    # Filtered view is cleanest without inline BU subtotals — Excel's
    # auto-filter dropdown lets the user slice by BU/Dept/Status, and the
    # grand-total row below uses SUBTOTAL so it tracks the visible rows.
    first_data_row = row
    for seq, emp in enumerate(employees, start=1):
        alloc_by_lt = emp.get("allocation_by_leave_type") or {}
        total_days = float(emp.get("total_days_allocated") or 0)

        values = [
            seq,
            emp.get("emp_code") or "—",
            emp.get("name") or "—",
            emp.get("email") or "—",
            emp.get("business_unit_name") or "—",
            emp.get("department_name") or "—",
            _format_label(emp.get("status_key") or ""),
            emp.get("date_of_joining") or "—",
            total_days,
        ]
        for lt_name in lt_names:
            values.append(float(alloc_by_lt.get(lt_name, 0) or 0))

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = _THIN_BORDER
            if col_idx == 1:
                cell.alignment = _CENTER_ALIGN
            elif col_idx >= 9:
                cell.alignment = _NUM_ALIGN
            else:
                cell.alignment = _LEFT_ALIGN
        row += 1

    last_data_row = row - 1

    # ── Auto-filter on header + data range ───────────────────────────────
    # Adds dropdown arrows on each header cell — users filter by BU, Dept,
    # Status, or any other column directly inside Excel.
    last_col_letter = get_column_letter(col_count)
    ws.auto_filter.ref = (
        f"A{header_row}:{last_col_letter}{last_data_row}"
        if last_data_row >= header_row
        else f"A{header_row}:{last_col_letter}{header_row}"
    )

    # ── Grand total (SUBTOTAL — updates with filter) ─────────────────────
    grand_vals: list = ["", "", "", "", "", "Grand Total", "", ""]
    if last_data_row >= first_data_row:
        for col_idx in range(9, col_count + 1):
            col = get_column_letter(col_idx)
            grand_vals.append(f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_data_row})")
    else:
        grand_vals.extend([0.0] * (col_count - 8))

    ws.row_dimensions[row].height = 24
    for col_idx, val in enumerate(grand_vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if col_idx >= 9 else _LEFT_ALIGN
    row += 1

    # ── Column widths ────────────────────────────────────────────────────
    widths = [6, 14, 26, 28, 22, 22, 16, 16, 20] + [16] * len(lt_names)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
