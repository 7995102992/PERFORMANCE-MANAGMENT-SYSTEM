import io
from datetime import date, datetime
from typing import Annotated, Any

from fastapi import APIRouter, Depends, Query, UploadFile, File, status
from fastapi.responses import StreamingResponse

from src.database import get_db_session
from src.dependencies import UserBase, get_current_user
from src.exceptions import DomainException
from src.holidays.schemas import (
    BulkHolidayFileImportRequest,
    BulkHolidayFileValidateResponse,
    BulkHolidayImport,
    BulkHolidayImportResponse,
    BulkHolidayItem,
    HolidayCreate,
    HolidayResponse,
    HolidayUpdate,
    PaginatedHolidayResponse,
    SyncHolidaysScopeRequest,
    SyncHolidaysScopeResponse,
)
from src.holidays.service import (
    bulk_import_holidays,
    create_holiday,
    delete_holiday,
    get_holidays_by_date_range,
    get_holidays_by_plan,
    sync_holidays_scope,
    update_holiday,
)
from src.leave_balance_processor.processor import run_daily_credit_processing
from src.logger import logger
from src.messaging import email_events
from src.utils import to_oid

router = APIRouter(tags=["holidays"])


async def _handle_action_flags(
    db,
    holiday_doc: dict,
    *,
    notify_employees: bool | None,
    reprocess_leaves: bool | None,
    holiday_id: str | None = None,
) -> None:
    """Execute post-create/update side-effects requested by the frontend.

    notify_employees — queues a holiday reminder email to every employee
        assigned to the holiday's plan via the transactional outbox.
    reprocess_leaves — re-runs the daily leave balance processor so that
        accruals are recalculated in light of the new/changed holiday.
    """
    if notify_employees:
        plan_id = holiday_doc.get("plan_id")
        if plan_id:
            plan = await db["holiday_plans"].find_one({"_id": plan_id, "deleted_on": None})
            plan_name = plan.get("name", "") if plan else ""

            member_docs = await db["holiday_plan_employees"].find(
                {"plan_id": plan_id, "deleted_on": None}
            ).to_list(length=None)
            user_ids = [d["user_id"] for d in member_docs]

            if user_ids:
                emp_docs = await db["employees"].find(
                    {"user_id": {"$in": user_ids}, "is_deleted": {"$ne": True}}
                ).to_list(length=None)

                holiday_name = holiday_doc.get("name", "")
                holiday_date = str(holiday_doc.get("date", ""))
                recipients = [e for e in emp_docs if e.get("work_email")]

                logger.info(
                    "Triggering holiday_reminder emails",
                    holiday_id=holiday_id,
                    holiday_name=holiday_name,
                    holiday_date=holiday_date,
                    plan_name=plan_name,
                    recipient_count=len(recipients),
                )
                for emp in recipients:
                    await email_events.publish_holiday_reminder(
                        employee_email=emp["work_email"],
                        employee_name=emp.get("name", ""),
                        holiday_name=holiday_name,
                        holiday_date=holiday_date,
                        plan_name=plan_name,
                        tenant_id=str(emp.get("organisation_id", "")),
                    )
        else:
            logger.warning("notify_employees=True but holiday has no plan_id", holiday_id=holiday_id)

    if reprocess_leaves:
        from datetime import datetime, timezone
        today = datetime.now(timezone.utc).date()
        logger.info("Triggering leave balance reprocessing", holiday_id=holiday_id, run_date=str(today))
        try:
            await run_daily_credit_processing(db, today)
            logger.info("Leave balance reprocessing completed", holiday_id=holiday_id)
        except Exception as exc:
            logger.error("Leave balance reprocessing failed", holiday_id=holiday_id, error=str(exc))


def _require_org_id(user: UserBase) -> str:
    if not user.org_id:
        raise DomainException(
            message="Authenticated user is missing an org_id",
            code="MISSING_ORG_CONTEXT",
            status_code=status.HTTP_403_FORBIDDEN,
        )
    return user.org_id


@router.post("/holidays", response_model=HolidayResponse, status_code=status.HTTP_201_CREATED)
async def add_holiday(
    payload: HolidayCreate,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> HolidayResponse:
    org_id = _require_org_id(current_user)
    doc = await create_holiday(db, payload, org_id=org_id, user_id=current_user.user_id)
    await _handle_action_flags(
        db,
        doc,
        notify_employees=payload.notify_employees,
        reprocess_leaves=payload.reprocess_leaves,
        holiday_id=str(doc.get("_id") or doc.get("id")),
    )
    return HolidayResponse(**doc)


@router.get("/holiday-plans/{plan_id}/holidays", response_model=PaginatedHolidayResponse)
async def list_holidays(
    plan_id: str,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    current_user: UserBase = Depends(get_current_user),
    db: Any = Depends(get_db_session),
) -> PaginatedHolidayResponse:
    items, total = await get_holidays_by_plan(
        db, plan_id, page, page_size,
        expected_org_id=current_user.org_id,
    )
    return PaginatedHolidayResponse(
        items=[HolidayResponse(**d) for d in items],
        total=total,
        page=page,
        page_size=page_size,
    )


@router.get("/holidays/calendar", response_model=list[HolidayResponse])
async def list_holidays_calendar(
    plan_id: str = Query(...),
    from_date: date = Query(...),
    to_date: date = Query(...),
    current_user: UserBase = Depends(get_current_user),
    db: Any = Depends(get_db_session),
) -> list[HolidayResponse]:
    org_id = _require_org_id(current_user)
    docs = await get_holidays_by_date_range(db, plan_id, org_id, from_date, to_date)
    return [HolidayResponse(**d) for d in docs]


@router.post("/holidays/bulk", response_model=BulkHolidayImportResponse)
async def bulk_import(
    payload: BulkHolidayImport,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> BulkHolidayImportResponse:
    org_id = _require_org_id(current_user)
    result = await bulk_import_holidays(
        db, payload.plan_id, payload.holidays, org_id=org_id, user_id=current_user.user_id
    )
    return BulkHolidayImportResponse(**result)


@router.put("/holidays/{holiday_id}", response_model=HolidayResponse)
async def modify_holiday(
    holiday_id: str,
    payload: HolidayUpdate,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> HolidayResponse:
    doc = await update_holiday(
        db, holiday_id, payload,
        user_id=current_user.user_id,
        current_user_org_id=current_user.org_id,
    )
    await _handle_action_flags(
        db,
        doc,
        notify_employees=payload.notify_employees,
        reprocess_leaves=payload.reprocess_leaves,
        holiday_id=holiday_id,
    )
    return HolidayResponse(**doc)


@router.delete("/holidays/{holiday_id}", status_code=status.HTTP_204_NO_CONTENT)
async def remove_holiday(
    holiday_id: str,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> None:
    await delete_holiday(
        db, holiday_id,
        user_id=current_user.user_id,
        current_user_org_id=current_user.org_id,
    )


@router.post(
    "/holiday-plans/{plan_id}/holidays/sync-scope",
    response_model=SyncHolidaysScopeResponse,
)
async def sync_holiday_scope(
    plan_id: str,
    payload: SyncHolidaysScopeRequest,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> SyncHolidaysScopeResponse:
    org_id = _require_org_id(current_user)
    result = await sync_holidays_scope(
        db,
        plan_id,
        extend=payload.extend,
        trim=payload.trim,
        bu_extend=payload.bu_extend,
        bu_trim=payload.bu_trim,
        org_id=org_id,
        user_id=current_user.user_id,
    )
    return SyncHolidaysScopeResponse(**result)


# ---------------------------------------------------------------------------
# Bulk holiday file-based import (XLSX/CSV template + validate + import)
# ---------------------------------------------------------------------------

async def _get_plan_for_bulk(db, plan_id: str, org_id: str) -> dict:
    """Resolve and validate a holiday plan for bulk operations."""
    plan_query: dict = {"_id": to_oid(plan_id), "deleted_on": None}
    plan_query["org_id"] = {"$in": [to_oid(org_id), str(org_id)]}
    plan = await db["holiday_plans"].find_one(plan_query)
    if not plan:
        raise DomainException(
            message="Holiday plan not found",
            code="PLAN_NOT_FOUND",
            status_code=status.HTTP_404_NOT_FOUND,
        )
    if plan.get("is_active") is False:
        raise DomainException(
            message="Cannot mutate holidays: the holiday plan is inactive.",
            code="INACTIVE_PLAN",
            status_code=status.HTTP_409_CONFLICT,
        )
    return plan


async def _get_classification_map(db, org_id: str) -> dict[str, dict]:
    """Return a mapping of lowercase classification name -> {_id, name, color}
    for the given org."""
    docs = await db["holiday_classifications"].find(
        {"org_id": {"$in": [to_oid(org_id), str(org_id)]}, "deleted_on": None}
    ).to_list(length=None)
    return {
        doc["name"].strip().lower(): {
            "_id": doc["_id"],
            "name": doc["name"],
            "color": doc.get("color", "#6b7280"),
        }
        for doc in docs
        if doc.get("name")
    }


def _generate_holiday_template_xlsx(
    plan_year: int,
    classification_names: list[str],
) -> bytes:
    """Generate an XLSX template for bulk holiday import."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Holidays"

    # Style constants (matching bulk_utils.py pattern)
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="366092")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    columns = [
        {"header": "Holiday Name", "width": 30},
        {"header": "Date", "width": 18},
        {"header": "Classification", "width": 25},
        {"header": "Description", "width": 40},
    ]

    # Header row
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    for col_idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col["header"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        # Use get_column_letter for columns beyond Z (future-safe)
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col_idx)].width = col["width"]

    # Add comments with instructions on each header
    ws.cell(row=1, column=1).comment = Comment(
        "Required. The name of the holiday (max 100 characters).",
        "System",
    )
    ws.cell(row=1, column=2).comment = Comment(
        "Required. Accepted formats:\n"
        f"  YYYY-MM-DD (e.g. {plan_year}-01-01)\n"
        f"  DD/MM/YYYY (e.g. 01/01/{plan_year})\n"
        f"  MM/DD/YYYY (e.g. 01/01/{plan_year})\n"
        f"  DD-MMM-YYYY (e.g. 01-Jan-{plan_year})\n"
        f"  DD-Mon-YY (e.g. 01-Jan-{str(plan_year)[-2:]})\n"
        f"The year is taken from the plan ({plan_year}). If you enter a date "
        f"with a different year, you can confirm an auto-fix during upload.",
        "System",
    )
    ws.cell(row=1, column=3).comment = Comment(
        "Required. Pick a classification from the dropdown.",
        "System",
    )
    ws.cell(row=1, column=4).comment = Comment(
        "Optional. A short description of the holiday.",
        "System",
    )

    # --- Classification dropdown (column C) ---
    if classification_names:
        ref = wb.create_sheet("_REF")
        ref.sheet_state = "hidden"
        for idx, name in enumerate(classification_names, start=1):
            ref.cell(row=idx, column=1, value=name)
        cls_dv = DataValidation(
            type="list",
            formula1=f"=_REF!$A$1:$A${len(classification_names)}",
            allow_blank=False,
            showErrorMessage=True,
        )
        cls_dv.error = "Pick a classification from the dropdown"
        cls_dv.errorTitle = "Invalid Classification"
        ws.add_data_validation(cls_dv)
        cls_dv.add("C2:C1001")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _coerce_cell_value(val, *, is_date_column: bool = False) -> str:
    """Convert an openpyxl cell value to a clean string.

    Handles datetime/date objects (→ YYYY-MM-DD) and avoids losing
    date information when Excel stores dates as serial numbers.
    Set ``is_date_column=True`` only for the Date column so that
    Excel serial numbers (e.g. 46030) are converted to dates without
    accidentally converting Year values like 2026.
    """
    if val is None:
        return ""
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, date):
        return val.strftime("%Y-%m-%d")
    if is_date_column and isinstance(val, (int, float)) and val > 30_000:
        try:
            from openpyxl.utils.datetime import from_excel
            return from_excel(val).strftime("%Y-%m-%d")
        except Exception:
            pass
    return str(val).strip()


def _parse_holiday_xlsx_rows(file_bytes: bytes) -> list[dict]:
    """Parse an XLSX file and return rows as dicts with normalised keys."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        return []

    # Normalise headers: strip whitespace, lowercase, replace spaces with underscores
    headers = []
    for h in raw_headers:
        if h is None:
            headers.append("")
        else:
            headers.append(str(h).strip().lower().replace(" ", "_"))

    # Map normalised header names to column indices
    col_map = {name: idx for idx, name in enumerate(headers) if name}

    result = []
    for row in rows_iter:
        # Skip completely empty rows
        if all(c is None or str(c).strip() == "" for c in row):
            continue
        d: dict = {}
        for name, idx in col_map.items():
            val = row[idx] if idx < len(row) else None
            d[name] = _coerce_cell_value(val, is_date_column=(name == "date"))
        result.append(d)
    return result


def _parse_holiday_csv_rows(file_bytes: bytes) -> list[dict]:
    """Parse a CSV file and return rows as dicts with normalised keys."""
    import csv

    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise DomainException(
            message="Unable to decode file as UTF-8. Re-save the file as UTF-8 CSV or upload an XLSX.",
            code="INVALID_FILE_ENCODING",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        return []

    # Build a normalised-to-original header mapping
    norm_map = {}
    for fn in reader.fieldnames:
        norm = fn.strip().lower().replace(" ", "_")
        norm_map[norm] = fn

    result = []
    for row in reader:
        if all(not v or not v.strip() for v in row.values()):
            continue
        d: dict = {}
        for norm, orig in norm_map.items():
            val = row.get(orig)
            d[norm] = val.strip() if val else ""
        result.append(d)
    return result


def _is_xlsx(file_bytes: bytes) -> bool:
    return file_bytes[:4] == b"PK\x03\x04"


def _resolve_header(row: dict, *candidates: str) -> str:
    """Return the value for the first matching key in row."""
    for c in candidates:
        v = row.get(c)
        if v is not None:
            return v
    return ""


def _parse_date_flexible(date_str: str) -> date | None:
    """Try multiple date formats and return a date object, or None if all fail.

    Supports:
    - YYYY-MM-DD (ISO)
    - YYYY-MM-DD HH:MM:SS (datetime with time suffix)
    - DD/MM/YYYY
    - MM/DD/YYYY
    - DD-MMM-YYYY (e.g. 01-Jan-2026)
    - DD-Mon-YY (e.g. 01-Jan-26)
    """
    if not date_str:
        return None

    # Strip leading/trailing whitespace
    date_str = date_str.strip()

    # Handle datetime strings with time suffix (e.g. "2026-01-01 00:00:00")
    clean_date = date_str.split(" ")[0] if " " in date_str else date_str

    # Try ISO format first: YYYY-MM-DD
    try:
        return datetime.strptime(clean_date, "%Y-%m-%d").date()
    except ValueError:
        pass

    # Try DD-MMM-YYYY (e.g. 01-Jan-2026)
    try:
        return datetime.strptime(clean_date, "%d-%b-%Y").date()
    except ValueError:
        pass

    # Try DD-Mon-YY (e.g. 01-Jan-26)
    try:
        return datetime.strptime(clean_date, "%d-%b-%y").date()
    except ValueError:
        pass

    # Try DD/MM/YYYY — but we also need to try MM/DD/YYYY.
    # Heuristic: if the first segment > 12, it must be DD/MM/YYYY.
    # If the second segment > 12, it must be MM/DD/YYYY.
    # If both <= 12, prefer DD/MM/YYYY (more common internationally).
    if "/" in clean_date:
        parts = clean_date.split("/")
        if len(parts) == 3:
            try:
                p1, p2, p3 = int(parts[0]), int(parts[1]), int(parts[2])
                if p1 > 12:
                    # Must be DD/MM/YYYY
                    return datetime.strptime(clean_date, "%d/%m/%Y").date()
                elif p2 > 12:
                    # Must be MM/DD/YYYY
                    return datetime.strptime(clean_date, "%m/%d/%Y").date()
                else:
                    # Ambiguous — prefer DD/MM/YYYY
                    return datetime.strptime(clean_date, "%d/%m/%Y").date()
            except (ValueError, IndexError):
                pass

    # Excel serial date number (e.g. "46030" for 2026-01-08)
    try:
        serial = int(float(clean_date))
        if 1 < serial < 200_000:
            from openpyxl.utils.datetime import from_excel
            return from_excel(serial).date()
    except (ValueError, TypeError):
        pass

    return None


async def _validate_holiday_rows(
    db,
    parsed_rows: list[dict],
    plan: dict,
    classification_map: dict[str, dict],
) -> dict:
    """Validate parsed holiday rows against the plan and classifications.

    Returns the shape expected by BulkHolidayFileValidateResponse.
    """
    plan_year = plan["year"]
    plan_oid = plan["_id"]

    # Fetch existing holiday dates for this plan to detect duplicates
    existing_dates: set[str] = set()
    cursor = db["holidays"].find(
        {"plan_id": plan_oid, "deleted_on": None},
        {"date": 1},
    )
    async for doc in cursor:
        existing_dates.add(str(doc["date"]))

    rows = []
    seen_dates: set[str] = set()
    valid_count = 0
    error_count = 0
    duplicate_count = 0
    year_mismatch_count = 0

    for i, parsed in enumerate(parsed_rows):
        row_num = i + 1
        errors: list[str] = []

        # --- Extract fields with flexible header matching ---
        name = _resolve_header(parsed, "holiday_name", "name", "holiday")
        date_str = _resolve_header(parsed, "date", "holiday_date")
        classification_str = _resolve_header(parsed, "classification", "classification_name", "type")
        description = _resolve_header(parsed, "description", "desc", "notes")

        # --- Validate name ---
        if not name:
            errors.append("Holiday Name is required")
        elif len(name) > 100:
            errors.append("Holiday Name must be at most 100 characters")

        # --- Validate date (multi-format) ---
        parsed_date = None
        if not date_str:
            errors.append("Date is required")
        else:
            parsed_date = _parse_date_flexible(date_str)
            if parsed_date is None:
                errors.append(
                    f"Invalid date format '{date_str}'. "
                    f"Accepted formats: YYYY-MM-DD, DD/MM/YYYY, MM/DD/YYYY, DD-MMM-YYYY, DD-Mon-YY"
                )
            else:
                date_str = parsed_date.isoformat()

        # --- Detect year mismatch against plan year (not a blocking error) ---
        # If the date's year doesn't match the plan year, surface this as a
        # `year_mismatch` row with a suggested corrected date the FE can confirm
        # before importing.
        year_mismatch = (
            parsed_date is not None and parsed_date.year != plan_year
        )
        suggested_date: str | None = None
        if year_mismatch and parsed_date:
            try:
                fixed = parsed_date.replace(year=plan_year)
            except ValueError:
                # Feb 29 on a non-leap plan year — fall back to Feb 28
                fixed = parsed_date.replace(year=plan_year, day=28)
            suggested_date = fixed.isoformat()

        # --- Validate classification ---
        matched_classification = None
        if not classification_str:
            errors.append("Classification is required")
        else:
            matched_classification = classification_map.get(classification_str.lower())
            if not matched_classification:
                valid_names = ", ".join(
                    c["name"] for c in classification_map.values()
                )
                errors.append(
                    f"Unknown classification '{classification_str}'. "
                    f"Valid options: {valid_names}"
                )

        # --- Build row result ---
        row_year_out = str(parsed_date.year) if parsed_date else None

        if errors:
            rows.append({
                "row_num": row_num,
                "name": name or None,
                "year": row_year_out,
                "date": date_str or None,
                "classification": classification_str or None,
                "classification_id": None,
                "description": description or None,
                "status": "error",
                "errors": errors,
                "suggested_date": None,
            })
            error_count += 1
            continue

        # Duplicate checks compare against the *effective* date — for
        # year_mismatch rows that's the suggested (plan-year) date, since
        # that's what will actually be imported if the user confirms.
        effective_date = suggested_date if year_mismatch else date_str

        # --- Check for duplicate date within the upload file ---
        if effective_date in seen_dates:
            rows.append({
                "row_num": row_num,
                "name": name,
                "year": row_year_out,
                "date": date_str,
                "classification": matched_classification["name"],
                "classification_id": str(matched_classification["_id"]),
                "description": description or None,
                "status": "error",
                "errors": [f"Duplicate date {effective_date} within this file"],
                "suggested_date": suggested_date,
            })
            error_count += 1
            continue

        # --- Check for duplicate date against existing holidays in plan ---
        if effective_date in existing_dates:
            rows.append({
                "row_num": row_num,
                "name": name,
                "year": row_year_out,
                "date": date_str,
                "classification": matched_classification["name"],
                "classification_id": str(matched_classification["_id"]),
                "description": description or None,
                "status": "duplicate",
                "errors": [f"A holiday already exists on {effective_date} in this plan"],
                "suggested_date": suggested_date,
            })
            duplicate_count += 1
            seen_dates.add(effective_date)
            continue

        seen_dates.add(effective_date)

        if year_mismatch:
            rows.append({
                "row_num": row_num,
                "name": name,
                "year": row_year_out,
                "date": date_str,
                "classification": matched_classification["name"],
                "classification_id": str(matched_classification["_id"]),
                "description": description or None,
                "status": "year_mismatch",
                "errors": [
                    f"Date year ({parsed_date.year}) does not match the plan year "
                    f"({plan_year}). Will be updated to {suggested_date} on confirm."
                ],
                "suggested_date": suggested_date,
            })
            year_mismatch_count += 1
            continue

        rows.append({
            "row_num": row_num,
            "name": name,
            "year": row_year_out,
            "date": date_str,
            "classification": matched_classification["name"],
            "classification_id": str(matched_classification["_id"]),
            "description": description or None,
            "status": "valid",
            "errors": [],
            "suggested_date": None,
        })
        valid_count += 1

    return {
        "rows": rows,
        "summary": {
            "total": len(parsed_rows),
            "valid": valid_count,
            "errors": error_count,
            "duplicates": duplicate_count,
            "year_mismatches": year_mismatch_count,
        },
    }


@router.get("/holiday-plans/{plan_id}/holidays/bulk-template")
async def download_holiday_bulk_template(
    plan_id: str,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
):
    """Generate and download an XLSX template for bulk holiday import."""
    org_id = _require_org_id(current_user)
    plan = await _get_plan_for_bulk(db, plan_id, org_id)
    classification_map = await _get_classification_map(db, org_id)

    # Sorted list of classification display names for the dropdown
    classification_names = sorted(c["name"] for c in classification_map.values())

    content = _generate_holiday_template_xlsx(plan["year"], classification_names)
    return StreamingResponse(
        io.BytesIO(content),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": (
                f'attachment; filename="holiday_template_{plan_id}_{plan["year"]}.xlsx"'
            ),
        },
    )


@router.post(
    "/holiday-plans/{plan_id}/holidays/bulk-validate",
    response_model=BulkHolidayFileValidateResponse,
)
async def validate_holiday_bulk_upload(
    plan_id: str,
    file: UploadFile = File(...),
    current_user: Annotated[UserBase, Depends(get_current_user)] = None,
    db: Any = Depends(get_db_session),
) -> BulkHolidayFileValidateResponse:
    """Accept an XLSX or CSV upload, parse and validate each row against the
    plan year, existing holidays, and classification names."""
    org_id = _require_org_id(current_user)
    plan = await _get_plan_for_bulk(db, plan_id, org_id)
    classification_map = await _get_classification_map(db, org_id)

    file_bytes = await file.read()

    if not file_bytes:
        raise DomainException(
            message="Uploaded file is empty",
            code="EMPTY_FILE",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    # Dispatch to XLSX or CSV parser
    if _is_xlsx(file_bytes):
        parsed_rows = _parse_holiday_xlsx_rows(file_bytes)
    else:
        parsed_rows = _parse_holiday_csv_rows(file_bytes)

    if not parsed_rows:
        return BulkHolidayFileValidateResponse(
            rows=[],
            summary={"total": 0, "valid": 0, "errors": 0, "duplicates": 0},
        )

    result = await _validate_holiday_rows(db, parsed_rows, plan, classification_map)
    return BulkHolidayFileValidateResponse(**result)


@router.post(
    "/holiday-plans/{plan_id}/holidays/bulk-import",
    response_model=BulkHolidayImportResponse,
)
async def bulk_import_from_file(
    plan_id: str,
    payload: BulkHolidayFileImportRequest,
    current_user: Annotated[UserBase, Depends(get_current_user)],
    db: Any = Depends(get_db_session),
) -> BulkHolidayImportResponse:
    """Accept validated holiday rows and insert them via the existing bulk
    import logic. Automatically applies the plan's BU/dept scope to each
    holiday so the caller doesn't need to provide them."""
    org_id = _require_org_id(current_user)
    plan = await _get_plan_for_bulk(db, plan_id, org_id)

    # Pull the plan's BU and dept scope so each imported holiday inherits it
    plan_bu_ids = [str(b) for b in (plan.get("business_unit_ids") or [])]
    plan_dept_ids = [str(d) for d in (plan.get("department_ids") or [])]

    if not plan_bu_ids:
        raise DomainException(
            message="Holiday plan has no business units configured. Cannot import holidays.",
            code="PLAN_MISSING_BU",
            status_code=status.HTTP_400_BAD_REQUEST,
        )
    if not plan_dept_ids:
        raise DomainException(
            message="Holiday plan has no departments configured. Cannot import holidays.",
            code="PLAN_MISSING_DEPT",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    # Convert the file-import items into BulkHolidayItem instances that the
    # existing bulk_import_holidays service function expects.
    bulk_items = []
    for h in payload.holidays:
        bulk_items.append(BulkHolidayItem(
            name=h.name,
            date=h.date,
            classification_id=h.classification_id,
            applicable_department_ids=plan_dept_ids,
            business_unit_ids=plan_bu_ids,
        ))

    result = await bulk_import_holidays(
        db, plan_id, bulk_items, org_id=org_id, user_id=current_user.user_id,
    )
    return BulkHolidayImportResponse(**result)
