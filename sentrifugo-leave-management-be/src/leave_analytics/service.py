"""
Leave analytics service.

Queries year-end execution records, enriches with employee and leave-type
metadata, and builds the data for the analytics table + XLSX export.
"""

import io
from typing import Optional

from motor.motor_asyncio import AsyncIOMotorDatabase
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.logger import logger
from src.utils import to_oid

EXECUTION_COLLECTION = "leave_year_end_execution"

# ─── Styles ──────────────────────────────────────────────────────────────────

_TITLE_FONT = Font(bold=True, size=14, color="1F3864")
_SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
_SECTION_FILL = PatternFill("solid", fgColor="366092")
_HEADER_FONT = Font(bold=True, size=10, color="1F3864")
_HEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUM_ALIGN = Alignment(horizontal="right", vertical="center")
_LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
_THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
)
_TOTAL_FONT = Font(bold=True, size=10, color="FFFFFF")


async def _get_leave_type_names(db: AsyncIOMotorDatabase, leave_plan_id: str) -> dict:
    plan = await db["leave_plans"].find_one(
        {"_id": to_oid(leave_plan_id), "deleted_on": None},
        {"leave_type_ids": 1},
    )
    leave_type_ids = list((plan or {}).get("leave_type_ids") or [])

    if not leave_type_ids:
        mappings = await db["leave_plan_type_mapping"].find(
            {"leave_plan_id": to_oid(leave_plan_id)},
            {"leave_type_id": 1},
        ).to_list(length=None)
        leave_type_ids = [m["leave_type_id"] for m in mappings]

    if not leave_type_ids:
        return {}

    lt_docs = await db["leave_types"].find(
        {"_id": {"$in": leave_type_ids}, "deleted_on": None},
    ).to_list(length=None)
    return {str(doc["_id"]): doc["name"] for doc in lt_docs}


async def _get_employee_details(db: AsyncIOMotorDatabase, employee_ids: list) -> dict:
    if not employee_ids:
        return {}

    oid_list = [to_oid(eid) for eid in employee_ids]
    emp_docs = await db["employees"].find(
        {"_id": {"$in": oid_list}},
        {
            "_id": 1, "emp_code": 1, "name": 1, "first_name": 1, "last_name": 1,
            "department_id": 1, "business_unit_id": 1,
        },
    ).to_list(length=None)

    dept_ids = {doc["department_id"] for doc in emp_docs if doc.get("department_id")}
    bu_ids = {doc["business_unit_id"] for doc in emp_docs if doc.get("business_unit_id")}

    dept_map = {}
    if dept_ids:
        dept_docs = await db["departments"].find({"_id": {"$in": list(dept_ids)}}).to_list(length=None)
        dept_map = {str(d["_id"]): d["name"] for d in dept_docs}

    bu_map = {}
    if bu_ids:
        bu_docs = await db["business_units"].find({"_id": {"$in": list(bu_ids)}}).to_list(length=None)
        bu_map = {str(d["_id"]): d["name"] for d in bu_docs}

    result = {}
    for emp in emp_docs:
        eid = str(emp["_id"])
        name = (
            emp.get("name")
            or f"{emp.get('first_name', '') or ''} {emp.get('last_name', '') or ''}".strip()
            or "—"
        )
        dept_id = str(emp["department_id"]) if emp.get("department_id") else None
        bu_id = str(emp["business_unit_id"]) if emp.get("business_unit_id") else None
        result[eid] = {
            "name": name,
            "emp_code": emp.get("emp_code") or "—",
            "department_name": dept_map.get(dept_id, "—") if dept_id else "—",
            "bu_name": bu_map.get(bu_id, "—") if bu_id else "—",
        }
    return result


async def get_year_end_report_data(
    db: AsyncIOMotorDatabase,
    leave_plan_id: Optional[str] = None,
    year: Optional[int] = None,
) -> dict:
    query: dict = {}
    if leave_plan_id:
        query["leave_plan_id"] = to_oid(leave_plan_id)
    if year:
        query["year"] = year

    execution_records = await db[EXECUTION_COLLECTION].find(query).to_list(length=None)

    if not execution_records:
        return {"records": [], "leave_types": [], "plans": [], "years": []}

    plan_ids = list({str(r["leave_plan_id"]) for r in execution_records})
    employee_ids = list({str(r["employee_id"]) for r in execution_records})

    employee_details = await _get_employee_details(db, employee_ids)

    all_leave_type_names: dict = {}
    for pid in plan_ids:
        lt_names = await _get_leave_type_names(db, pid)
        all_leave_type_names.update(lt_names)

    plan_docs = await db["leave_plans"].find(
        {"_id": {"$in": [to_oid(pid) for pid in plan_ids]}},
        {"_id": 1, "name": 1},
    ).to_list(length=None)
    plan_name_map = {str(d["_id"]): d["name"] for d in plan_docs}

    years = sorted({r["year"] for r in execution_records}, reverse=True)

    lt_ids = list(all_leave_type_names.keys())

    records = []
    for r in execution_records:
        emp_id = str(r.get("employee_id", ""))
        emp = employee_details.get(emp_id, {})
        per_type = r.get("per_type_details") or {}

        per_type_data = {}
        for lt_id in lt_ids:
            lt_name = all_leave_type_names[lt_id]
            detail = per_type.get(lt_id, {})
            per_type_data[lt_name] = {
                "opening_balance": detail.get("opening_balance", 0),
                "payout": detail.get("payout", 0),
                "carry_forward": detail.get("carry_forward", 0),
                "expired": detail.get("expired", 0),
                "closing_balance": detail.get("closing_balance", 0),
            }

        records.append({
            "id": str(r["_id"]),
            "employee_id": emp_id,
            "emp_code": emp.get("emp_code", "—"),
            "employee_name": emp.get("name", "—"),
            "department": emp.get("department_name", "—"),
            "business_unit": emp.get("bu_name", "—"),
            "leave_plan_id": str(r["leave_plan_id"]),
            "leave_plan_name": plan_name_map.get(str(r["leave_plan_id"]), "—"),
            "year": r["year"],
            "opening_balance": r.get("opening_balance", 0),
            "payout_amount": r.get("payout_amount", 0),
            "carry_forward_amount": r.get("carry_forward_amount", 0),
            "expired_amount": r.get("expired_amount", 0),
            "negative_recovered_amount": r.get("negative_recovered_amount", 0),
            "closing_balance": r.get("closing_balance", 0),
            "per_type": per_type_data,
            "execution_status": r.get("execution_status", ""),
            "processed_at": r.get("processed_at").isoformat() if r.get("processed_at") else None,
        })

    records.sort(key=lambda x: (x["business_unit"], x["department"], x["employee_name"]))

    plans = [{"id": pid, "name": plan_name_map.get(pid, "—")} for pid in plan_ids]

    return {
        "records": records,
        "leave_types": [{"id": k, "name": v} for k, v in all_leave_type_names.items()],
        "plans": plans,
        "years": years,
    }


def generate_year_end_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Year-End Report"

    records = data.get("records", [])
    leave_types = data.get("leave_types", [])
    lt_names = [lt["name"] for lt in leave_types]

    base_headers = ["#", "Emp Code", "Employee Name", "Department", "Business Unit", "Leave Plan", "Year"]

    per_type_headers = []
    for lt_name in lt_names:
        per_type_headers.extend([
            f"{lt_name} - Opening",
            f"{lt_name} - Payout",
            f"{lt_name} - Carry Fwd",
            f"{lt_name} - Expired",
            f"{lt_name} - Closing",
        ])

    summary_headers = ["Total Opening", "Total Payout", "Total Carry Fwd", "Total Expired", "Total Closing", "Status"]
    all_headers = base_headers + per_type_headers + summary_headers
    col_count = len(all_headers)

    num_start = len(base_headers) + 1
    num_end = col_count - 1

    row = 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value="Year-End Processing Report")
    cell.font = _TITLE_FONT
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 28
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    banner = ws.cell(row=row, column=1, value="Employee Year-End Breakdown by Leave Type")
    banner.font = _SECTION_FONT
    banner.fill = _SECTION_FILL
    banner.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 24
    row += 1

    ws.row_dimensions[row].height = 30
    for col_idx, header in enumerate(all_headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    header_row = row
    row += 1

    first_data_row = row
    for seq, rec in enumerate(records, start=1):
        values = [
            seq,
            rec.get("emp_code", "—"),
            rec.get("employee_name", "—"),
            rec.get("department", "—"),
            rec.get("business_unit", "—"),
            rec.get("leave_plan_name", "—"),
            rec.get("year", ""),
        ]

        per_type = rec.get("per_type", {})
        for lt_name in lt_names:
            detail = per_type.get(lt_name, {})
            values.extend([
                detail.get("opening_balance", 0),
                detail.get("payout", 0),
                detail.get("carry_forward", 0),
                detail.get("expired", 0),
                detail.get("closing_balance", 0),
            ])

        values.extend([
            rec.get("opening_balance", 0),
            rec.get("payout_amount", 0),
            rec.get("carry_forward_amount", 0),
            rec.get("expired_amount", 0),
            rec.get("closing_balance", 0),
            rec.get("execution_status", ""),
        ])

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.border = _THIN_BORDER
            if col_idx == 1:
                cell.alignment = _CENTER_ALIGN
            elif num_start <= col_idx <= num_end:
                cell.alignment = _NUM_ALIGN
            else:
                cell.alignment = _LEFT_ALIGN
        row += 1

    last_data_row = row - 1

    last_col = get_column_letter(col_count)
    ws.auto_filter.ref = f"A{header_row}:{last_col}{max(last_data_row, header_row)}"

    grand_vals = ["", "", "", "", "", "", "Grand Total"]
    if last_data_row >= first_data_row:
        for col_idx in range(num_start, num_end + 1):
            col = get_column_letter(col_idx)
            grand_vals.append(f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_data_row})")
    else:
        grand_vals.extend([0.0] * (num_end - num_start + 1))
    grand_vals.append("")

    ws.row_dimensions[row].height = 24
    for col_idx, val in enumerate(grand_vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = _TOTAL_FONT
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if num_start <= col_idx <= num_end else _LEFT_ALIGN

    widths = [6, 14, 26, 22, 22, 22, 8]
    widths.extend([14] * len(per_type_headers))
    widths.extend([16, 14, 16, 14, 16, 12])
    for i, w in enumerate(widths, start=1):
        if i <= col_count:
            ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=len(base_headers) + 1).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
