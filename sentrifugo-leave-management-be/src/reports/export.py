import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_TITLE_FONT = Font(bold=True, size=14, color="1F3864")
_SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
_SECTION_FILL = PatternFill("solid", fgColor="366092")
_HEADER_FONT = Font(bold=True, size=10, color="1F3864")
_HEADER_FILL = PatternFill("solid", fgColor="D6E4F0")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUM_ALIGN = Alignment(horizontal="right", vertical="center")
_LEFT_ALIGN = Alignment(horizontal="left", vertical="center")
_CENTER_ALIGN = Alignment(horizontal="center", vertical="center")
_THIN_BORDER = Border(
    bottom=Side(style="thin", color="D9D9D9"),
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
)
_TOTAL_FONT = Font(bold=True, size=10)
_TOTAL_FILL = PatternFill("solid", fgColor="F2F2F2")


def _set_cell(ws, row, col, value, *, font=None, fill=None, alignment=None):
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if alignment:
        cell.alignment = alignment
    elif isinstance(value, (int, float)):
        cell.alignment = _NUM_ALIGN
    else:
        cell.alignment = _LEFT_ALIGN
    cell.border = _THIN_BORDER
    return cell


def generate_year_end_report_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Year-End Processing Report"

    leave_types = data.get("leave_types", [])
    lt_names = [lt["name"] for lt in leave_types]

    base_headers = ["#", "Emp Code", "Employee Name", "Department", "Business Unit"]

    per_type_headers = []
    for lt_name in lt_names:
        for suffix in ["Opening", "Payout", "Encash", "Carry Fwd", "Expired", "Closing"]:
            per_type_headers.append(f"{lt_name} - {suffix}")

    total_headers = ["Total Opening", "Total Payout", "Total Encash", "Total Carry Fwd", "Total Expired", "Total Closing", "Status"]
    all_headers = base_headers + per_type_headers + total_headers
    col_count = len(all_headers)

    row = 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value="Year-End Processing Report")
    cell.font = _TITLE_FONT
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 28
    row += 1

    plan_name = data.get("leave_plan_name", "")
    year = data.get("year", "")
    ws.cell(row=row, column=1, value="Leave Plan").font = Font(bold=True, size=10)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row=row, column=2, value=plan_name)
    row += 1
    ws.cell(row=row, column=1, value="Year").font = Font(bold=True, size=10)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row=row, column=2, value=str(year))
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value="Employee Year-End Breakdown by Leave Type")
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 24
    row += 1

    ws.row_dimensions[row].height = 30
    for col_idx, header in enumerate(all_headers, start=1):
        _set_cell(ws, row, col_idx, header, font=_HEADER_FONT, fill=_HEADER_FILL, alignment=_HEADER_ALIGN)
    header_row = row
    row += 1

    employees = data.get("employees", [])
    first_data_row = row
    for seq, emp in enumerate(employees, start=1):
        details = emp.get("leave_type_details", [])
        detail_map = {d["leave_type_name"]: d for d in details}

        values = [
            seq,
            emp.get("emp_code", ""),
            emp.get("name", ""),
            emp.get("department", ""),
            emp.get("business_unit", ""),
        ]

        for lt_name in lt_names:
            d = detail_map.get(lt_name, {})
            values.append(d.get("opening_balance", 0))
            values.append(d.get("payout_amount", 0))
            values.append(d.get("encashed_amount", 0))
            values.append(d.get("carry_forward_amount", 0))
            values.append(d.get("expired_amount", 0))
            values.append(d.get("closing_balance", 0))

        values.extend([
            emp.get("opening_balance", 0),
            emp.get("payout_amount", 0),
            emp.get("encashed_amount", 0),
            emp.get("carry_forward_amount", 0),
            emp.get("expired_amount", 0),
            emp.get("closing_balance", 0),
            emp.get("status", "SUCCESS"),
        ])

        for col_idx, val in enumerate(values, start=1):
            if col_idx == 1:
                _set_cell(ws, row, col_idx, val, alignment=_CENTER_ALIGN)
            else:
                _set_cell(ws, row, col_idx, val)
        row += 1

    last_data_row = row - 1

    last_col_letter = get_column_letter(col_count)
    if last_data_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"

    grand_vals = ["", "", "", "", "Grand Total"]
    # 6 per-type numeric columns (Opening/Payout/Encash/Carry/Expired/Closing) ×
    # types, plus the 6 total numeric columns; the trailing Status column is not
    # summed (the "" appended below).
    num_cols = len(per_type_headers) + 6
    if last_data_row >= first_data_row:
        for col_idx in range(6, 6 + num_cols):
            col = get_column_letter(col_idx)
            grand_vals.append(f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_data_row})")
    else:
        grand_vals.extend([0] * num_cols)
    grand_vals.append("")

    ws.row_dimensions[row].height = 24
    for col_idx, val in enumerate(grand_vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if col_idx > 5 else _LEFT_ALIGN

    # 5 base + (6 per type) + 6 total numeric (incl. Encash) + Status.
    widths = [6, 14, 26, 22, 22] + [14] * len(per_type_headers) + [16, 16, 16, 16, 16, 16, 12]
    for i, w in enumerate(widths, start=1):
        if i <= col_count:
            ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=6).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_current_balance_xlsx(data: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Current Leave Balance"

    leave_types = data.get("leave_types", [])
    lt_names = [lt["name"] for lt in leave_types]

    per_type_headers = []
    for lt_name in lt_names:
        for suffix in ["Opening", "Used YTD", "Balance"]:
            per_type_headers.append(f"{lt_name} - {suffix}")

    base_headers = ["#", "Emp Code", "Employee Name", "Department", "Business Unit"]
    total_headers = ["Total Opening", "Total Used YTD", "Total Balance"]
    all_headers = base_headers + per_type_headers + total_headers
    col_count = len(all_headers)

    row = 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value="Current Leave Balance Report")
    cell.font = _TITLE_FONT
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 28
    row += 1

    plan_name = data.get("leave_plan_name", "")
    ws.cell(row=row, column=1, value="Leave Plan").font = Font(bold=True, size=10)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row=row, column=2, value=plan_name)
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value="Employee Current Leave Balances")
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 24
    row += 1

    ws.row_dimensions[row].height = 30
    for col_idx, header in enumerate(all_headers, start=1):
        _set_cell(ws, row, col_idx, header, font=_HEADER_FONT, fill=_HEADER_FILL, alignment=_HEADER_ALIGN)
    header_row = row
    row += 1

    employees = data.get("employees", [])
    first_data_row = row
    for seq, emp in enumerate(employees, start=1):
        details = emp.get("leave_type_details", [])
        detail_map = {d["leave_type_name"]: d for d in details}

        values = [
            seq,
            emp.get("emp_code", ""),
            emp.get("name", ""),
            emp.get("department", ""),
            emp.get("business_unit", ""),
        ]
        for lt_name in lt_names:
            d = detail_map.get(lt_name, {})
            values.append(d.get("opening", 0))
            values.append(d.get("used_ytd", 0))
            values.append(d.get("balance", 0))

        values.extend([
            emp.get("opening", 0),
            emp.get("used_ytd", 0),
            emp.get("total_balance", 0),
        ])

        for col_idx, val in enumerate(values, start=1):
            if col_idx == 1:
                _set_cell(ws, row, col_idx, val, alignment=_CENTER_ALIGN)
            else:
                _set_cell(ws, row, col_idx, val)
        row += 1

    last_data_row = row - 1

    last_col_letter = get_column_letter(col_count)
    if last_data_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{last_data_row}"

    grand_vals = ["", "", "", "", "Grand Total"]
    num_cols = len(per_type_headers) + 3
    if last_data_row >= first_data_row:
        for col_idx in range(6, 6 + num_cols):
            col = get_column_letter(col_idx)
            grand_vals.append(f"=SUBTOTAL(9,{col}{first_data_row}:{col}{last_data_row})")
    else:
        grand_vals.extend([0] * num_cols)

    ws.row_dimensions[row].height = 24
    for col_idx, val in enumerate(grand_vals, start=1):
        cell = ws.cell(row=row, column=col_idx, value=val)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if col_idx > 5 else _LEFT_ALIGN

    widths = [6, 14, 26, 22, 22] + [14] * len(per_type_headers) + [16, 16, 16]
    for i, w in enumerate(widths, start=1):
        if i <= col_count:
            ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = ws.cell(row=header_row + 1, column=6).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── Employee leave report ────────────────────────────────────────────────────
# A single sheet: the request-level extract HR filters on screen, one row per
# leave request, carrying the approval that closed it. The filters the run used
# are stamped into the header so a saved file still says what it is later.

_LEAVE_DATA_HEADERS = [
    "#", "Emp Code", "Employee Name", "Business Unit", "Department", "Designation",
    "Leave Type", "From Date", "To Date", "Days", "Duration", "Status", "Loss of Pay",
    "Applied On", "L1 Manager", "L2 Manager", "Action By", "Action On",
    "Turnaround (Days)", "Pending With", "Reason", "Approver Comment",
]

_LEAVE_DATA_WIDTHS = [
    6, 14, 26, 22, 22, 22, 18, 13, 13, 9, 16, 12, 12, 14, 22, 22, 22, 14, 16, 22, 34, 34,
]

# 1-indexed position of the Days column in _LEAVE_DATA_HEADERS — the only column
# a grand total means anything for.
_DAYS_COL = 10


def _duration_label(row: dict) -> str:
    mode = (row.get("duration_mode") or "").upper()
    if mode == "HALF_DAY":
        half = (row.get("half_day_period") or "").replace("_", " ").title()
        return f"Half Day ({half})" if half else "Half Day"
    if mode == "CUSTOM":
        return "Custom"
    return "Full Day(s)"


def _title_block(ws, row: int, col_count: int, title: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = _TITLE_FONT
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 28
    return row + 1


def _meta_row(ws, row: int, label: str, value) -> int:
    ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)
    ws.cell(row=row, column=2, value=value)
    return row + 1


def _section_banner(ws, row: int, col_count: int, text: str) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_count)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _SECTION_FONT
    cell.fill = _SECTION_FILL
    cell.alignment = _LEFT_ALIGN
    ws.row_dimensions[row].height = 24
    return row + 1


def _header_row(ws, row: int, headers: list) -> int:
    ws.row_dimensions[row].height = 30
    for col_idx, header in enumerate(headers, start=1):
        _set_cell(ws, row, col_idx, header, font=_HEADER_FONT, fill=_HEADER_FILL, alignment=_HEADER_ALIGN)
    return row + 1


def generate_employee_leave_xlsx(data: dict) -> bytes:
    """Workbook for the HR employee-leave report — the request-level extract only.

    Deliberately a single sheet: the statistics are a read-on-screen summary, not
    something the download carries. Everything needed to recompute them is in
    these rows, and Excel's own filters and pivots do it better than a frozen
    second sheet would.
    """
    rows = data.get("rows", [])
    filters = data.get("filters_applied", {})

    wb = Workbook()
    ws = wb.active
    ws.title = "Leave Data"
    col_count = len(_LEAVE_DATA_HEADERS)
    period = f"{data.get('from_date', '')} to {data.get('to_date', '')}"

    row = _title_block(ws, 1, col_count, "Employee Leave Report")
    row = _meta_row(ws, row, "Period", period)
    for label, value in filters.items():
        row = _meta_row(ws, row, label, value)
    if data.get("truncated"):
        row = _meta_row(
            ws, row, "Note",
            f"Truncated to the first {data.get('max_rows')} rows — "
            "narrow the filters for a complete extract.",
        )
    row += 1

    row = _section_banner(ws, row, col_count, "Leave Requests & Approvals")
    header_row = row
    row = _header_row(ws, row, _LEAVE_DATA_HEADERS)

    first_data_row = row
    for seq, r in enumerate(rows, start=1):
        turnaround = r.get("approval_turnaround_days")
        values = [
            seq,
            r.get("emp_code", ""),
            r.get("employee_name", ""),
            r.get("business_unit", ""),
            r.get("department", ""),
            r.get("designation", ""),
            r.get("leave_type", ""),
            r.get("from_date", ""),
            r.get("to_date", ""),
            r.get("days", 0),
            _duration_label(r),
            r.get("status", ""),
            "Yes" if r.get("loss_of_pay") else "No",
            (r.get("applied_on") or "")[:10],
            r.get("l1_manager", "—"),
            r.get("l2_manager", "—"),
            r.get("action_by") or "—",
            (r.get("action_on") or "")[:10] or "—",
            turnaround if turnaround is not None else "—",
            r.get("current_approver") or "—",
            r.get("reason", ""),
            r.get("action_comment") or "",
        ]
        for col_idx, value in enumerate(values, start=1):
            if col_idx == 1:
                _set_cell(ws, row, col_idx, value, alignment=_CENTER_ALIGN)
            else:
                _set_cell(ws, row, col_idx, value)
        row += 1

    last_data_row = row - 1
    if last_data_row >= header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(col_count)}{last_data_row}"

    # Total the Days column only — every other numeric column is a per-request
    # measure that would be meaningless summed. SUBTOTAL so it follows the filter.
    days_letter = get_column_letter(_DAYS_COL)
    total_values = [""] * col_count
    total_values[4] = "Grand Total"
    total_values[_DAYS_COL - 1] = (
        f"=SUBTOTAL(9,{days_letter}{first_data_row}:{days_letter}{last_data_row})"
        if last_data_row >= first_data_row else 0
    )
    ws.row_dimensions[row].height = 24
    for col_idx, value in enumerate(total_values, start=1):
        cell = ws.cell(row=row, column=col_idx, value=value)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = _SECTION_FILL
        cell.border = _THIN_BORDER
        cell.alignment = _NUM_ALIGN if col_idx == _DAYS_COL else _LEFT_ALIGN

    for i, width in enumerate(_LEAVE_DATA_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=4).coordinate

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
