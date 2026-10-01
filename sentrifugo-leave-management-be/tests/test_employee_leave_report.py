"""Unit tests for the HR employee-leave report.

Covers the pure shaping layer — row assembly from a request + its approval
trail, the statistics rollup, and the workbook generator — which is where the
reporting logic actually lives. The Mongo-facing helpers are thin `find` calls
and are exercised by the API tests instead.
"""

import io
from datetime import date
from unittest.mock import MagicMock

import pytest
from bson import ObjectId
from openpyxl import load_workbook

from src.exceptions import DomainException
from src.reports.employee_leave import (
    _as_date_str,
    _build_row,
    _build_statistics,
    _day_bounds,
    _oid_list,
    _requested,
    _resolve_employees,
    _summarise,
    _validate_range,
)
from src.reports.export import _duration_label, generate_employee_leave_xlsx

LEAVE_TYPES = {"lt1": "Earned Leave", "lt2": "Sick Leave"}
MANAGERS = {"m1": "Ravi Kumar", "m2": "Priya Nair"}

EMPLOYEE = {
    "employee_id": "e1",
    "user_id": "u1",
    "emp_code": "SIL-001",
    "employee_name": "Asha Rao",
    "email": "asha@example.com",
    "department": "Engineering",
    "business_unit": "Products",
    "designation": "SDE II",
    "l1_manager_id": "m1",
    "l2_manager_id": "m2",
}


def _request(**overrides):
    base = {
        "_id": "r1",
        "user_id": "u1",
        "leave_type_id": "lt1",
        "start_date": "2026-03-02",
        "end_date": "2026-03-04",
        "duration_days": 3.0,
        "duration_hours": 24.0,
        "duration_mode": "FULL_DAYS",
        "status": "APPROVED",
        "loss_of_pay": False,
        "reason": "Family function",
        "approval_state": {"current_level": 1},
    }
    base.update(overrides)
    return base


def _trail(*entries):
    return list(entries)


SUBMITTED = {
    "action": "SUBMITTED", "level": None, "actor_id": "u1",
    "actor_name": "Asha Rao", "comment": None,
    "acted_on": "2026-02-20T09:00:00+00:00",
}
APPROVED_L1 = {
    "action": "APPROVED", "level": 1, "actor_id": "m1",
    "actor_name": "Ravi Kumar", "comment": "Approved",
    "acted_on": "2026-02-22T09:00:00+00:00",
}
REJECTED_L2 = {
    "action": "REJECTED", "level": 2, "actor_id": "m2",
    "actor_name": "Priya Nair", "comment": "Team crunch",
    "acted_on": "2026-02-23T09:00:00+00:00",
}


# ── Date handling ────────────────────────────────────────────────────────────

def test_validate_range_rejects_reversed_window():
    with pytest.raises(DomainException) as exc:
        _validate_range(date(2026, 3, 31), date(2026, 3, 1))
    assert exc.value.code == "INVALID_DATE_RANGE"


def test_validate_range_allows_single_day():
    _validate_range(date(2026, 3, 1), date(2026, 3, 1))


def test_day_bounds_span_the_whole_end_day():
    """A leave starting late on to_date must still fall inside the window."""
    start, end = _day_bounds(date(2026, 3, 1), date(2026, 3, 31))
    assert start.day == 1 and start.hour == 0
    assert end.day == 31 and end.hour == 23 and end.minute == 59


@pytest.mark.parametrize(
    "raw,expected",
    [
        ("2026-03-02", "2026-03-02"),
        ("2026-03-02 00:00:00", "2026-03-02"),
        ("2026-03-02T10:30:00+00:00", "2026-03-02"),
        (date(2026, 3, 2), "2026-03-02"),
        (None, None),
    ],
)
def test_as_date_str_normalises_legacy_shapes(raw, expected):
    assert _as_date_str(raw) == expected


# ── Row assembly ─────────────────────────────────────────────────────────────

def test_approved_row_carries_the_final_approver_and_turnaround():
    row = _build_row(
        _request(), EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED, APPROVED_L1), MANAGERS
    )
    assert row["action_by"] == "Ravi Kumar"
    assert row["action_comment"] == "Approved"
    assert row["approval_turnaround_days"] == 2.0
    assert row["current_approver"] is None
    assert row["l1_manager"] == "Ravi Kumar"
    assert row["l2_manager"] == "Priya Nair"
    assert row["leave_type"] == "Earned Leave"


def test_last_decision_wins_when_a_request_passes_two_levels():
    row = _build_row(
        _request(status="REJECTED"),
        EMPLOYEE, LEAVE_TYPES,
        _trail(SUBMITTED, APPROVED_L1, REJECTED_L2),
        MANAGERS,
    )
    assert row["action_by"] == "Priya Nair"
    assert row["action_comment"] == "Team crunch"


def test_pending_row_names_who_it_is_waiting_on():
    row = _build_row(
        _request(status="PENDING"), EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED), MANAGERS
    )
    assert row["current_approver"] == "Ravi Kumar"
    assert row["action_by"] is None
    assert row["approval_turnaround_days"] is None


def test_pending_at_level_two_points_at_the_l2_manager():
    row = _build_row(
        _request(status="PENDING", approval_state={"current_level": 2}),
        EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED), MANAGERS,
    )
    assert row["current_approver"] == "Priya Nair"


def test_pending_falls_back_to_l1_when_no_l2_is_set():
    employee = {**EMPLOYEE, "l2_manager_id": None}
    row = _build_row(
        _request(status="PENDING", approval_state={"current_level": 2}),
        employee, LEAVE_TYPES, _trail(SUBMITTED), MANAGERS,
    )
    assert row["current_approver"] == "Ravi Kumar"
    assert row["l2_manager"] == "—"


def test_row_survives_a_request_with_no_activity_at_all():
    row = _build_row(_request(), EMPLOYEE, LEAVE_TYPES, [], MANAGERS)
    assert row["action_by"] is None
    assert row["approval_trail"] == []
    assert row["days"] == 3.0


def test_unknown_leave_type_does_not_blank_the_row():
    row = _build_row(_request(leave_type_id="gone"), EMPLOYEE, LEAVE_TYPES, [], MANAGERS)
    assert row["leave_type"] == "—"
    assert row["leave_type_id"] == "gone"


# ── Statistics ───────────────────────────────────────────────────────────────

def _rows_for_stats():
    approved = _build_row(
        _request(), EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED, APPROVED_L1), MANAGERS
    )
    pending = _build_row(
        _request(_id="r2", status="PENDING", duration_days=1.0, loss_of_pay=True),
        EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED), MANAGERS,
    )
    other_bu = _build_row(
        _request(_id="r3", leave_type_id="lt2", duration_days=2.0),
        {**EMPLOYEE, "user_id": "u2", "employee_name": "Dev Shah",
         "business_unit": "Services", "department": "Support"},
        LEAVE_TYPES, _trail(SUBMITTED, APPROVED_L1), MANAGERS,
    )
    return [approved, pending, other_bu]


def test_statistics_split_by_status_and_count_days():
    stats = _build_statistics(_rows_for_stats(), 42, date(2026, 3, 1), date(2026, 3, 31))
    assert stats["total_requests"] == 3
    assert stats["total_days"] == 6.0
    assert stats["approved_requests"] == 2
    assert stats["approved_days"] == 5.0
    assert stats["pending_requests"] == 1
    assert stats["pending_days"] == 1.0
    assert stats["rejected_requests"] == 0
    assert stats["employees_in_scope"] == 42


def test_only_approved_leave_counts_an_employee_as_on_leave():
    """A pending request is not time off yet — it must not inflate the headline."""
    pending_only = [
        _build_row(
            _request(status="PENDING"), EMPLOYEE, LEAVE_TYPES, _trail(SUBMITTED), MANAGERS
        )
    ]
    stats = _build_statistics(pending_only, 10, date(2026, 3, 1), date(2026, 3, 31))
    assert stats["employees_on_leave"] == 0

    stats_mixed = _build_statistics(_rows_for_stats(), 42, date(2026, 3, 1), date(2026, 3, 31))
    # u1 (approved) and u2 (approved) — u1's pending row must not double-count.
    assert stats_mixed["employees_on_leave"] == 2


def test_loss_of_pay_days_are_tracked_separately():
    stats = _build_statistics(_rows_for_stats(), 42, date(2026, 3, 1), date(2026, 3, 31))
    assert stats["loss_of_pay_days"] == 1.0


def test_turnaround_average_ignores_still_pending_requests():
    stats = _build_statistics(_rows_for_stats(), 42, date(2026, 3, 1), date(2026, 3, 31))
    # Both approved rows took exactly 2 days; the pending one contributes nothing.
    assert stats["avg_approval_turnaround_days"] == 2.0


def test_breakdowns_group_and_rank_by_days():
    stats = _build_statistics(_rows_for_stats(), 42, date(2026, 3, 1), date(2026, 3, 31))
    assert stats["by_leave_type"] == [
        {"name": "Earned Leave", "requests": 2, "days": 4.0},
        {"name": "Sick Leave", "requests": 1, "days": 2.0},
    ]
    assert stats["by_business_unit"] == [
        {"name": "Products", "requests": 2, "days": 4.0},
        {"name": "Services", "requests": 1, "days": 2.0},
    ]
    assert [d["name"] for d in stats["by_department"]] == ["Engineering", "Support"]


def test_empty_result_set_yields_zeroed_statistics():
    stats = _build_statistics([], 0, date(2026, 3, 1), date(2026, 3, 31))
    assert stats["total_requests"] == 0
    assert stats["total_days"] == 0
    assert stats["avg_days_per_request"] == 0.0
    assert stats["avg_approval_turnaround_days"] is None
    assert stats["by_leave_type"] == []
    assert stats["from_date"] == "2026-03-01"


def test_summarise_buckets_missing_values_under_a_dash():
    assert _summarise([{"business_unit": None, "days": 1.0}], "business_unit") == [
        {"name": "—", "requests": 1, "days": 1.0}
    ]


# ── Workbook ─────────────────────────────────────────────────────────────────

@pytest.mark.parametrize(
    "row,expected",
    [
        ({"duration_mode": "FULL_DAYS"}, "Full Day(s)"),
        ({"duration_mode": "HALF_DAY", "half_day_period": "FIRST_HALF"}, "Half Day (First Half)"),
        ({"duration_mode": "HALF_DAY", "half_day_period": None}, "Half Day"),
        ({"duration_mode": "CUSTOM"}, "Custom"),
        ({}, "Full Day(s)"),
    ],
)
def test_duration_label(row, expected):
    assert _duration_label(row) == expected


def _workbook(rows, **extra):
    payload = {
        "from_date": "2026-03-01", "to_date": "2026-03-31",
        "rows": rows, "truncated": False, "max_rows": 20000, "filters_applied": {},
    }
    payload.update(extra)
    return load_workbook(io.BytesIO(generate_employee_leave_xlsx(payload)))


def test_workbook_is_the_extract_only_with_one_row_per_request():
    rows = _rows_for_stats()
    wb = _workbook(rows, filters_applied={"Business Unit": "Products"})

    # Extract only — the summary is an on-screen view, not part of the download.
    assert wb.sheetnames == ["Leave Data"]
    sheet = wb["Leave Data"]
    header_row = next(
        r for r in range(1, sheet.max_row + 1) if sheet.cell(row=r, column=1).value == "#"
    )
    names = [
        sheet.cell(row=header_row + i, column=3).value for i in range(1, len(rows) + 1)
    ]
    assert names == [r["employee_name"] for r in rows]
    # Grand-total banner sits directly below the last data row.
    assert sheet.cell(row=header_row + len(rows) + 1, column=5).value == "Grand Total"


def test_workbook_needs_no_statistics_key():
    """The generator must not reach for a summary the export no longer builds."""
    wb = _workbook(_rows_for_stats())
    assert wb.sheetnames == ["Leave Data"]


def test_workbook_records_the_filters_it_was_run_with():
    sheet = _workbook(
        _rows_for_stats(), filters_applied={"Employees": "Asha Rao", "Status": "APPROVED"}
    )["Leave Data"]
    labels = [sheet.cell(row=r, column=1).value for r in range(1, 12)]
    assert "Employees" in labels and "Status" in labels


def test_workbook_renders_with_no_rows():
    assert _workbook([]).sheetnames == ["Leave Data"]


def test_truncation_note_is_written_into_the_sheet():
    sheet = _workbook([], truncated=True, max_rows=5)["Leave Data"]
    labels = [sheet.cell(row=r, column=1).value for r in range(1, 12)]
    assert "Note" in labels


# ── Employee filter composition ──────────────────────────────────────────────
# The employee picker and the free-text search both need $or. A naive
# implementation lets one overwrite the other, which silently WIDENS the report
# — the worst possible failure for something HR downloads and acts on.

OID_A = "507f1f77bcf86cd799439011"
OID_B = "507f1f77bcf86cd799439012"


def _capture_employee_match(**kwargs) -> dict:
    """Run _resolve_employees against a stub db and return the employees filter."""
    captured = {}

    def _find(match, _projection):
        captured["match"] = match
        cursor = MagicMock()

        async def _to_list(length=None):
            return []

        cursor.to_list = _to_list
        return cursor

    db = MagicMock()
    db.__getitem__.return_value.find = _find

    import asyncio

    asyncio.run(_resolve_employees(db, None, **kwargs))
    return captured["match"]


def test_employee_selection_matches_either_id_form():
    match = _capture_employee_match(employee_ids=[OID_A])
    clause = match["$and"][0]["$or"]
    assert {"_id": {"$in": [ObjectId(OID_A)]}} in clause
    # user_id is stored as ObjectId or string depending on the writer.
    assert {"user_id": {"$in": [ObjectId(OID_A), OID_A]}} in clause


def test_employee_selection_and_search_both_survive():
    match = _capture_employee_match(employee_ids=[OID_A], search="asha")
    assert len(match["$and"]) == 2, "one clause overwrote the other"
    assert "$or" not in match, "top-level $or would let a later clause clobber it"
    joined = str(match["$and"])
    assert OID_A in joined and "asha" in joined


def test_bu_and_department_filters_coexist_with_an_employee_selection():
    match = _capture_employee_match(
        business_unit_ids=[OID_A], department_ids=[OID_B], employee_ids=[OID_A]
    )
    assert "business_unit_id" in match and "department_id" in match
    assert len(match["$and"]) == 1


def test_no_filters_leaves_the_match_unconstrained():
    match = _capture_employee_match()
    assert "$and" not in match
    assert match == {"is_deleted": {"$ne": True}}


def test_an_all_invalid_employee_selection_matches_nobody():
    """Never fall back to "everyone" when the requested ids are unparseable."""
    match = _capture_employee_match(employee_ids=["not-an-oid"])
    clause = match["$and"][0]["$or"]
    assert clause == [{"_id": {"$in": []}}, {"user_id": {"$in": []}}]


def test_oid_list_skips_malformed_entries_without_raising():
    assert _oid_list([OID_A, "garbage", "", None, OID_B]) == [
        ObjectId(OID_A), ObjectId(OID_B)
    ]


@pytest.mark.parametrize(
    "values,expected",
    [(None, False), ([], False), ([""], False), (["x"], True), (["", "x"], True)],
)
def test_requested_distinguishes_omitted_from_empty(values, expected):
    assert _requested(values) is expected
