"""Dump the leave history of one or more employees, in one org, to an Excel workbook.

Read-only. Given an org id and any number of employee identifiers, this exports
every leave request each person raised on or after a cut-off date (1 January of
the current year by default) — one worksheet per employee, in the order the ids
were passed.

Each row is one leave request, carrying its dates, duration, status, reason and
the approval that closed it. The approval columns come from
``leave_request_activity``: the finalizing approver is the LAST APPROVED entry,
so a multi-level flow reports the approval that actually closed the request
rather than level 1's.

Balances, holds and ledgers are deliberately not exported — this is the history
sheet only. Pass --output with a .json suffix instead to get the raw records for
scripting against.

Each employee identifier is matched, in order, against employees._id, then
user_id, then emp_code — so you can pass whichever id you have to hand. The
resolved employee's organisation_id must equal the org id you passed; an id from
another tenant comes back as "not found" rather than being exported.

An id that resolves to nobody is reported and skipped; the run still writes the
employees that did resolve, and exits non-zero so a caller notices the gap.

Usage:
    python -m scripts.export_employee_leave_data <org_id> <employee_id> [<employee_id> ...]
    python -m scripts.export_employee_leave_data <org_id> SIL-0856 SIL-0860 --output team.xlsx
    python -m scripts.export_employee_leave_data <org_id> <employee_id> --since 2025-04-01
    python -m scripts.export_employee_leave_data <org_id> <employee_id> --output emp.json
    python -m scripts.export_employee_leave_data <org_id> <employee_id> --csv-dir ./out
    python -m scripts.export_employee_leave_data <org_id> <employee_id> --quiet
"""

import argparse
import asyncio
import csv
import json
import os
import re
from datetime import date, datetime
from decimal import Decimal

from bson import ObjectId
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import settings
from src.utils import to_oid, uid_match

EMPLOYEES_COL = "employees"
REQUESTS_COL = "leave_requests"
ACTIVITY_COL = "leave_request_activity"
LEAVE_TYPES_COL = "leave_types"

OID_RE = re.compile(r"^[0-9a-fA-F]{24}$")

# The one shape this script writes: a leave-history row. The workbook sheets and
# the CSVs both take their columns from here, so they can never drift apart.
# emp_code/employee_name lead so several sheets can be stacked into one table.
HISTORY_FIELDS = [
    "emp_code", "employee_name", "request_id", "leave_type_code",
    "leave_type_name", "start_date", "end_date", "duration_mode",
    "duration_hours", "duration_days", "loss_of_pay", "status", "reason",
    "created_on", "approved_by", "approved_on",
]

# Cell typing for the workbook. Dates and numbers go into Excel as real dates
# and floats rather than as text — text dates sort lexically, and Excel would
# otherwise re-read an ISO date through the machine's own locale.
DATE_COLUMNS = {"start_date", "end_date"}
DATETIME_COLUMNS = {"created_on", "approved_on"}
NUMBER_COLUMNS = {"duration_hours", "duration_days"}

HEADER_FILL = PatternFill("solid", start_color="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _jsonable(value):
    """Make a Mongo document JSON-serialisable without dropping anything.

    ObjectId/datetime/Decimal stringify or narrow to float; everything else
    passes through unchanged.
    """
    if isinstance(value, ObjectId):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if type(value).__name__ == "Decimal128":  # bson.decimal128, when present
        return float(value.to_decimal())
    if isinstance(value, dict):
        return {k: _jsonable(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_jsonable(v) for v in value]
    return value


def _person_name(person: dict | None) -> str:
    if not person:
        return ""
    first = (person.get("first_name") or "").strip()
    last = (person.get("last_name") or "").strip()
    if first or last:
        return f"{first} {last}".strip()
    return person.get("full_name") or person.get("name") or ""


async def _resolve_employee(db, org_id: str, employee_id: str) -> dict | None:
    """Find the employee by _id, then user_id, then emp_code — scoped to the org.

    emp_code is only unique within an org, so it is always matched together with
    organisation_id. The two ObjectId lookups are org-scoped too, which is what
    turns a cross-tenant id into "not found" instead of a wrong-tenant export.
    """
    org_oid = to_oid(org_id)
    org_clause = {"organisation_id": {"$in": [org_oid, str(org_oid)]}}

    if OID_RE.match(employee_id):
        oid = to_oid(employee_id)
        for query in (
            {"_id": oid, **org_clause},
            {"user_id": uid_match(oid), **org_clause},
        ):
            doc = await db[EMPLOYEES_COL].find_one(query)
            if doc:
                return doc

    return await db[EMPLOYEES_COL].find_one({"emp_code": employee_id, **org_clause})


async def _leave_type_index(db, org_id: str) -> dict:
    """_id -> leave type doc, for every type visible to the org.

    Not filtered on deleted_on: historical requests can point at a type that has
    since been soft-deleted, and an export that dropped those labels would
    misreport what the leave was for.
    """
    org_oid = to_oid(org_id)
    types = await db[LEAVE_TYPES_COL].find(
        {"$or": [{"org_id": {"$in": [org_oid, str(org_oid)]}}, {"org_id": None}]}
    ).to_list(length=None)
    return {t["_id"]: t for t in types}


def _type_label(types: dict, leave_type_id) -> tuple[str, str]:
    lt = types.get(to_oid(str(leave_type_id))) if leave_type_id else None
    if not lt:
        return "", ""
    return lt.get("code") or "", lt.get("name") or ""


def _row_start_date(request: dict) -> date | None:
    """The request's start date, whichever field carries it.

    start_date is an ISO string in this collection, but older rows can carry
    only start_datetime — read both so the cut-off never silently drops a
    request it should have kept.
    """
    raw = request.get("start_date")
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    if raw:
        try:
            return date.fromisoformat(str(raw)[:10])
        except ValueError:
            pass
    dt = request.get("start_datetime")
    if isinstance(dt, datetime):
        return dt.date()
    return None


async def collect_employee(db, org_id: str, employee_id: str, since: date,
                           types: dict) -> dict:
    """One employee's leave history from `since` onwards, plus who they are."""
    emp = await _resolve_employee(db, org_id, employee_id)
    if not emp:
        raise LookupError(
            f"No employee matching {employee_id!r} in org {org_id} "
            f"(tried _id, user_id, emp_code)."
        )

    user_oid = emp.get("user_id")
    emp_code = emp.get("emp_code") or ""
    display_name = _person_name(emp)

    requests: list[dict] = []
    if user_oid:
        # Filtered in Python rather than in the query: start_date is an ISO
        # string here but a few rows carry only start_datetime, and one Mongo
        # clause cannot compare both shapes. A row with neither is kept, so an
        # odd record shows up rather than vanishing.
        for r in await db[REQUESTS_COL].find(
            {"user_id": uid_match(user_oid)}
        ).sort("start_date", 1).to_list(length=None):
            start = _row_start_date(r)
            if start is None or start >= since:
                requests.append(r)

    activity_by_request: dict = {}
    actor_names: dict = {}
    if requests:
        acts = await db[ACTIVITY_COL].find(
            {"leave_request_id": {"$in": [r["_id"] for r in requests]}}
        ).sort("timestamp", 1).to_list(length=None)
        for act in acts:
            activity_by_request.setdefault(act["leave_request_id"], []).append(act)

        actor_ids = [a["actor_id"] for a in acts if a.get("actor_id")]
        if actor_ids:
            for e in await db[EMPLOYEES_COL].find(
                {"user_id": {"$in": actor_ids}},
                {"user_id": 1, "first_name": 1, "last_name": 1, "full_name": 1, "name": 1},
            ).to_list(length=None):
                actor_names[e["user_id"]] = _person_name(e)

    rows = []
    for r in requests:
        code, name = _type_label(types, r.get("leave_type_id"))
        trail = activity_by_request.get(r["_id"], [])
        # The finalizing approval is the LAST APPROVED entry, so a multi-level
        # flow reports the approval that closed the request, not level 1's.
        approvals = [a for a in trail if a.get("action") == "APPROVED"]
        final = approvals[-1] if approvals else None
        rows.append({
            **_jsonable(r),
            "emp_code": emp_code,
            "employee_name": display_name,
            "request_id": str(r["_id"]),
            "leave_type_code": code,
            "leave_type_name": name,
            "approved_by": (
                actor_names.get(final["actor_id"]) or str(final["actor_id"]) if final else ""
            ),
            "approved_on": _jsonable(final["timestamp"]) if final else "",
        })

    return {
        "query_id": employee_id,
        "employee_id": str(emp["_id"]),
        "user_id": str(user_oid) if user_oid else None,
        "emp_code": emp_code,
        "display_name": display_name,
        "leave_requests": rows,
    }


async def collect(db, org_id: str, employee_ids: list[str], since: date) -> dict:
    """Every requested employee's history. One bad id does not stop the rest."""
    types = await _leave_type_index(db, org_id)
    employees, errors = [], []
    for employee_id in employee_ids:
        try:
            employees.append(await collect_employee(db, org_id, employee_id, since, types))
        except LookupError as exc:
            errors.append(str(exc))
    return {
        "exported_on": datetime.now().astimezone().isoformat(),
        "database": db.name,
        "query": {
            "org_id": str(org_id),
            "employee_ids": employee_ids,
            "since": since.isoformat(),
        },
        "employees": employees,
        "errors": errors,
    }


def _coerce_cell(field: str, value):
    """Turn an exported (already stringified) value back into a typed cell.

    collect() renders dates and ObjectIds as strings so the JSON is faithful;
    Excel wants the dates and numbers back as real values. Anything that does
    not parse is left as-is rather than blanked — a surprising value should
    still show up in the sheet.
    """
    if value is None or value == "":
        return None
    if field in NUMBER_COLUMNS:
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    if field in DATE_COLUMNS or field in DATETIME_COLUMNS:
        if isinstance(value, (datetime, date)):
            parsed = value
        else:
            try:
                parsed = datetime.fromisoformat(str(value))
            except ValueError:
                return value
        if isinstance(parsed, datetime):
            # openpyxl cannot write tz-aware datetimes; the timestamps are UTC.
            return parsed.replace(tzinfo=None) if parsed.tzinfo else parsed
        return parsed
    if isinstance(value, (dict, list)):
        return json.dumps(value)
    return value


def _write_sheet(ws, rows: list[dict]) -> None:
    ws.append(HISTORY_FIELDS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([_coerce_cell(f, row.get(f)) for f in HISTORY_FIELDS])

    for idx, field in enumerate(HISTORY_FIELDS, start=1):
        letter = get_column_letter(idx)
        if field in DATE_COLUMNS or field in DATETIME_COLUMNS or field in NUMBER_COLUMNS:
            fmt = (
                "DD-MM-YYYY" if field in DATE_COLUMNS
                else "DD-MM-YYYY HH:MM" if field in DATETIME_COLUMNS
                else "0.00##"
            )
            for cell in ws[letter][1:]:  # skip the header
                cell.number_format = fmt
        widest = max(
            [len(field)] + [len(str(row.get(field) or "")) for row in rows]
        ) if rows else len(field)
        ws.column_dimensions[letter].width = min(widest + 4, 45)

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HISTORY_FIELDS))}{ws.max_row}"


def _sheet_title(emp: dict, taken: set) -> str:
    """A readable, unique, Excel-legal sheet name for one employee.

    Excel caps sheet names at 31 characters and rejects a handful of
    punctuation; two employees can also share a code-less name, so a collision
    takes a numeric suffix instead of overwriting the earlier sheet.
    """
    base = emp["emp_code"] or emp["display_name"] or emp["query_id"]
    base = re.sub(r"[\[\]:*?/\\]", "-", str(base)).strip() or str(emp["query_id"])
    title = base[:31]
    n = 2
    while title.lower() in taken:
        suffix = f" ({n})"
        title = base[:31 - len(suffix)] + suffix
        n += 1
    taken.add(title.lower())
    return title


def write_xlsx(data: dict, output_path: str) -> None:
    """One workbook, one leave-history sheet per employee, in the order passed."""
    wb = Workbook()
    wb.remove(wb.active)
    taken: set = set()
    for emp in data["employees"]:
        _write_sheet(wb.create_sheet(_sheet_title(emp, taken)), emp["leave_requests"])
    if not wb.sheetnames:  # every id failed to resolve — a workbook needs a sheet
        _write_sheet(wb.create_sheet("Leave History"), [])
    wb.save(output_path)


def write_csvs(data: dict, csv_dir: str) -> list[str]:
    os.makedirs(csv_dir, exist_ok=True)
    written = []
    for emp in data["employees"]:
        stem = re.sub(r"[^A-Za-z0-9_.-]", "_", emp["emp_code"] or emp["query_id"])
        path = os.path.join(csv_dir, f"{stem}_leave_history.csv")
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=HISTORY_FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(emp["leave_requests"])
        written.append(path)
    return written


def print_summary(data: dict) -> None:
    print(f"database:  {data['database']}")
    print(f"org_id:    {data['query']['org_id']}")
    print(f"since:     {data['query']['since']}")
    print()
    for emp in data["employees"]:
        print(f"  {(emp['emp_code'] or emp['query_id']):<16} "
              f"{(emp['display_name'] or '(no name)'):<28} "
              f"requests={len(emp['leave_requests'])}")
        # No user_id means no leave requests can exist — worth saying out loud,
        # otherwise an empty sheet reads as "took no leave".
        if not emp["user_id"]:
            print("    !! no user_id — this employee can have no leave requests")
    for err in data["errors"]:
        print(f"  !! {err}")


async def run(org_id: str, employee_ids: list[str], since: date, output: str,
              csv_dir: str | None, quiet: bool) -> int:
    if not settings.MONGODB_URL:
        print("ERROR: MONGODB_URL is not configured. Check your .env file.")
        return 1

    client = AsyncIOMotorClient(settings.MONGODB_URL)
    try:
        db = client.get_default_database()
        data = await collect(db, org_id, employee_ids, since)

        if os.path.splitext(output)[1].lower() == ".json":
            with open(output, "w", encoding="utf-8") as f:
                json.dump(data, f, indent=2, default=str)
        else:
            write_xlsx(data, output)

        if not quiet:
            print_summary(data)
            print()
        print(f"Wrote {output}")

        if csv_dir:
            for path in write_csvs(data, csv_dir):
                print(f"Wrote {path}")

        # Non-zero if any id failed to resolve — the file for the others was
        # still written, but a caller should not treat the run as complete.
        return 1 if data["errors"] else 0
    finally:
        client.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("org_id", help="organisations._id the employees belong to")
    parser.add_argument(
        "employee_ids", nargs="+",
        help="one or more employees._id, user_id, or emp_code — matched in that "
             "order; one worksheet each",
    )
    parser.add_argument(
        "--since", default=None,
        help="Only requests starting on or after this ISO date "
             "(default: 1 January of the current year)",
    )
    parser.add_argument(
        "--output", default=None,
        help="Output path; .xlsx writes the Excel workbook, .json writes the raw "
             "records (default: leave_history_<employee_id>.xlsx for one "
             "employee, leave_history.xlsx for several)",
    )
    parser.add_argument(
        "--csv-dir", default=None,
        help="Also write one flat CSV per employee into this directory",
    )
    parser.add_argument("--quiet", action="store_true", help="Skip the console summary")
    args = parser.parse_args()

    if args.since:
        try:
            since = date.fromisoformat(args.since)
        except ValueError:
            parser.error(f"--since must be an ISO date (YYYY-MM-DD), got {args.since!r}")
    else:
        since = date(date.today().year, 1, 1)

    output = args.output or (
        f"leave_history_{args.employee_ids[0]}.xlsx" if len(args.employee_ids) == 1
        else "leave_history.xlsx"
    )
    ext = os.path.splitext(output)[1].lower()
    if ext not in (".xlsx", ".json"):
        parser.error(f"--output must end in .xlsx or .json (got {ext or 'no extension'!r})")

    raise SystemExit(asyncio.run(
        run(args.org_id, args.employee_ids, since, output, args.csv_dir, args.quiet)
    ))


if __name__ == "__main__":
    main()
