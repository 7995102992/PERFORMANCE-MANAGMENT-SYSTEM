"""Export one organisation's leave requests in a date range to Excel or CSV.

Columns (dates in dd-mm-yyyy):

    Employee ID, Employee Name, Request Type, From Date, To Date,
    Duration (Days), Approval Status

"Duration (Days)" is the request's charged duration (0.5 for a half day,
weekends/holidays already excluded), not the raw span between the two dates.

The output format follows the --output extension: .xlsx (default) writes a
formatted worksheet, .csv writes plain CSV.

leave_requests carries no org field, so the org scope is resolved through
``employees.organisation_id`` -> ``user_id``, the same way the wipe/repair
scripts do it. Soft-deleted employees are included: their historical leave
still belongs in the org's export.

Usage:
    python -m scripts.export_leave_requests_by_date_range --org-id <oid> --from 01-04-2026 --to 30-06-2026
    ... --output q1_leaves.csv      # CSV instead of Excel
    ... --date-basis start          # only requests STARTING in the window
    ... --status APPROVED,PENDING   # filter by approval status
"""

import argparse
import asyncio
import csv
import os
from datetime import date, datetime

from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.config import settings
from src.utils import to_oid

EMPLOYEES_COL = "employees"
REQUESTS_COL = "leave_requests"
LEAVE_TYPES_COL = "leave_types"

FIELDNAMES = [
    "Employee ID",
    "Employee Name",
    "Request Type",
    "From Date",
    "To Date",
    "Duration (Days)",
    "Approval Status",
]

# Columns holding a real date / number, so the Excel writer can type them
# instead of dumping strings Excel would re-interpret with its own locale.
DATE_COLUMNS = {"From Date", "To Date"}
NUMBER_COLUMNS = {"Duration (Days)"}


def _parse_ddmmyyyy(value: str) -> date:
    try:
        return datetime.strptime(value.strip(), "%d-%m-%Y").date()
    except ValueError:
        raise argparse.ArgumentTypeError(
            f"{value!r} is not a valid dd-mm-yyyy date (e.g. 01-04-2026)"
        )


def _to_date(value) -> date | None:
    """Coerce a stored start/end date to a ``date``.

    Writers store ISO strings ("2026-04-01"), but a legacy row may hold a
    datetime or a bare date — accept all three rather than dropping the row.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _fmt(d: date | None) -> str:
    return d.strftime("%d-%m-%Y") if d else ""


def _name_of(emp: dict) -> str:
    name = emp.get("name")
    if name:
        return name
    return f"{emp.get('first_name') or ''} {emp.get('last_name') or ''}".strip()


def _duration_days(request: dict) -> float | None:
    """Leave days charged for a request — the applied duration, not the span.

    ``duration_days`` is what the service writes (0.5 for a half day, weekends
    and holidays already excluded unless the request is LOP), so it is the
    number that matches the balance movement. A row missing it is derived from
    ``duration_hours`` against the 8-hour workday the service uses.
    """
    days = request.get("duration_days")
    if days is None:
        hours = request.get("duration_hours")
        if hours is None:
            return None
        days = hours / 8.0
    try:
        return round(float(days), 2)
    except (TypeError, ValueError):
        return None


def _write_csv(rows: list[dict], output_path: str) -> None:
    with open(output_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        for row in rows:
            writer.writerow({
                k: (_fmt(v) if k in DATE_COLUMNS else ("" if v is None else v))
                for k, v in row.items()
            })


def _write_xlsx(rows: list[dict], output_path: str, sheet_title: str) -> None:
    """Write a formatted worksheet: frozen filterable header, typed cells.

    Dates go in as real dates carrying a dd-mm-yyyy display format rather than
    as text — text dates sort lexically and Excel would otherwise re-read
    "01-04-2026" through the machine's own locale.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]  # Excel caps sheet names at 31 chars

    ws.append(FIELDNAMES)
    header_fill = PatternFill("solid", start_color="1F4E79")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row[name] for name in FIELDNAMES])

    for idx, name in enumerate(FIELDNAMES, start=1):
        letter = get_column_letter(idx)
        if name in DATE_COLUMNS or name in NUMBER_COLUMNS:
            fmt = "DD-MM-YYYY" if name in DATE_COLUMNS else "0.0#"
            for cell in ws[letter][1:]:  # skip the header
                cell.number_format = fmt
                cell.alignment = Alignment(horizontal="center")
        widest = max(
            [len(name)] + [len(_fmt(r[name]) if name in DATE_COLUMNS else str(r[name] or "")) for r in rows]
        ) if rows else len(name)
        ws.column_dimensions[letter].width = min(widest + 4, 40)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(FIELDNAMES))}{ws.max_row}"
    wb.save(output_path)


async def export(
    org_id: str,
    from_date: date,
    to_date: date,
    output_path: str,
    date_basis: str,
    statuses: list[str] | None,
) -> None:
    if not settings.MONGODB_URL:
        print("ERROR: MONGODB_URL is not configured. Check your .env file.")
        return

    client = AsyncIOMotorClient(settings.MONGODB_URL)
    db = client.get_default_database()
    print(f"Database: {db.name}")

    try:
        employees = await db[EMPLOYEES_COL].find(
            {"organisation_id": to_oid(org_id)},
            {"user_id": 1, "emp_code": 1, "name": 1, "first_name": 1, "last_name": 1},
        ).to_list(length=None)

        if not employees:
            print(f"No employees found for organisation {org_id}. Nothing to export.")
            return

        # user_id has been stored as ObjectId (canonical) and as string
        # (legacy) — match both so no employee's requests are missed.
        emp_by_user_id: dict[str, dict] = {}
        user_id_match: list = []
        for e in employees:
            uid = e.get("user_id")
            if not uid:
                continue
            emp_by_user_id[str(uid)] = e
            user_id_match.extend([to_oid(str(uid)), str(uid)])

        print(
            f"Organisation {org_id}: {len(employees):,} employee(s), "
            f"{len(emp_by_user_id):,} with a user_id"
        )

        query: dict = {"user_id": {"$in": user_id_match}, "deleted_on": None}
        if statuses:
            query["status"] = {"$in": statuses}

        # Range filter as ISO strings (how the dates are stored) so Mongo does
        # the bulk of the work; a row stored as some other type still passes
        # through and is judged in Python below.
        lo, hi = from_date.isoformat(), to_date.isoformat()
        if date_basis == "start":
            range_clause = {"start_date": {"$gte": lo, "$lte": hi}}
        else:  # overlap
            range_clause = {"start_date": {"$lte": hi}, "end_date": {"$gte": lo}}
        query["$or"] = [range_clause, {"start_date": {"$not": {"$type": "string"}}}]

        requests = await db[REQUESTS_COL].find(query).to_list(length=None)

        type_ids = {r.get("leave_type_id") for r in requests if r.get("leave_type_id")}
        leave_types = await db[LEAVE_TYPES_COL].find(
            {"_id": {"$in": list(type_ids)}}, {"name": 1, "code": 1}
        ).to_list(length=None)
        type_by_id = {
            str(lt["_id"]): (lt.get("name") or lt.get("code") or "") for lt in leave_types
        }

        rows = []
        for r in requests:
            start = _to_date(r.get("start_date"))
            if start is None:
                continue
            end = _to_date(r.get("end_date")) or start

            # Re-apply the window in Python: authoritative for the rows the
            # string comparison could not judge, a no-op for the rest.
            if date_basis == "start":
                if not (from_date <= start <= to_date):
                    continue
            elif not (start <= to_date and end >= from_date):
                continue

            emp = emp_by_user_id.get(str(r.get("user_id")), {})
            rows.append({
                "Employee ID": emp.get("emp_code") or "",
                "Employee Name": _name_of(emp),
                "Request Type": type_by_id.get(str(r.get("leave_type_id")), ""),
                "From Date": start,
                "To Date": end,
                "Duration (Days)": _duration_days(r),
                "Approval Status": r.get("status") or "",
            })

        # Group each employee's rows together, oldest leave first.
        rows.sort(key=lambda row: (row["Employee Name"].lower(), row["From Date"]))

        if output_path.lower().endswith(".csv"):
            _write_csv(rows, output_path)
        else:
            _write_xlsx(rows, output_path, sheet_title="Leave Requests")

        print(
            f"Range {_fmt(from_date)} .. {_fmt(to_date)} ({date_basis} basis)"
            + (f", status in {statuses}" if statuses else "")
        )
        print(f"Wrote {len(rows):,} row(s) to {output_path}")
    finally:
        client.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "--org-id", required=True, help="organisation_id to export (required)"
    )
    parser.add_argument(
        "--from", dest="from_date", required=True, type=_parse_ddmmyyyy,
        help="Range start, dd-mm-yyyy (inclusive)",
    )
    parser.add_argument(
        "--to", dest="to_date", required=True, type=_parse_ddmmyyyy,
        help="Range end, dd-mm-yyyy (inclusive)",
    )
    parser.add_argument(
        "--output", default="leave_requests_export.xlsx",
        help="Output path; .xlsx writes Excel, .csv writes CSV "
             "(default: leave_requests_export.xlsx)",
    )
    parser.add_argument(
        "--date-basis", choices=["overlap", "start"], default="overlap",
        help="'overlap' (default) includes any request touching the range; "
             "'start' only those starting inside it",
    )
    parser.add_argument(
        "--status", default=None,
        help="Comma-separated approval statuses to include "
             "(PENDING,APPROVED,REJECTED,CANCELLED). Default: all",
    )
    args = parser.parse_args()

    if args.from_date > args.to_date:
        parser.error("--from must not be after --to")

    ext = os.path.splitext(args.output)[1].lower()
    if ext not in (".xlsx", ".csv"):
        parser.error(f"--output must end in .xlsx or .csv (got {ext or 'no extension'!r})")

    statuses = (
        [s.strip().upper() for s in args.status.split(",") if s.strip()]
        if args.status else None
    )
    asyncio.run(export(
        args.org_id, args.from_date, args.to_date,
        args.output, args.date_basis, statuses,
    ))


if __name__ == "__main__":
    main()
